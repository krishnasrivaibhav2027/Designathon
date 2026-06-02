from fastapi import APIRouter, HTTPException, status, Depends, UploadFile, File
from typing import List, Optional
from datetime import datetime
from app.core.database import get_db
from app.core.security import get_current_user, has_role, check_batch_access, check_report_card_access
from app.models.models import row_to_api
from app.services.pool_cleanup import clean_and_sync_pool
from app.schemas.report_card_schemas import (
    SparkReportCardResponse, SparkReportCardUpdate,
    FoundationReportCardResponse, FoundationReportCardUpdate,
    StreamReportCardResponse, StreamReportCardUpdate
)
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from fastapi.responses import StreamingResponse
import io
import csv

router = APIRouter(prefix="/api/report-card", tags=["report-card"])

# Helper: Convert snake_case Supabase keys to camelCase keys for API responses
def to_camel(snake_str):
    components = snake_str.split('_')
    return components[0] + ''.join(x.title() for x in components[1:])

def map_db_to_api(row: dict) -> dict:
    if not row:
        return {}
    res = {}
    for k, v in row.items():
        # Handle specific custom manual overrides
        if k == "ga1_a1": res["ga1_a1"] = v
        elif k == "ga1_a2": res["ga1_a2"] = v
        elif k == "ga2_a1": res["ga2_a1"] = v
        elif k == "ga2_a2": res["ga2_a2"] = v
        elif k == "ga3_a1": res["ga3_a1"] = v
        elif k == "ga3_a2": res["ga3_a2"] = v
        elif k == "ga4_a1": res["ga4_a1"] = v
        elif k == "ga4_a2": res["ga4_a2"] = v
        elif k == "ga5_a1": res["ga5_a1"] = v
        elif k == "ga5_a2": res["ga5_a2"] = v
        elif k == "project_eval_a1": res["projectEval_a1"] = v
        elif k == "project_eval_a2": res["projectEval_a2"] = v
        elif k == "final_grade_a1": res["finalGrade_a1"] = v
        elif k == "final_grade_a2": res["finalGrade_a2"] = v
        elif k == "contact_number": res["contactNumber"] = v
        elif k == "foundation_language": res["foundationLanguage"] = v
        elif k == "email_sent_date": res["emailSentDate"] = v
        elif k == "training_status": res["trainingStatus"] = v
        elif k == "superset_id": res["supersetId"] = v
        elif k == "batch_id": res["batchId"] = v
        elif k == "candidate_id": res["candidateId"] = v
        elif k == "created_at": res["createdAt"] = v
        elif k == "updated_at": res["updatedAt"] = v
        elif k == "trainer_name": res["trainerName"] = v
        elif k == "training_start_date": res["trainingStartDate"] = v
        elif k == "training_end_date": res["trainingEndDate"] = v
        elif k == "batch_no": res["batchNo"] = v
        elif k == "a1_score": res["a1Score"] = v
        elif k == "a2_score": res["a2Score"] = v
        elif k == "final_status": res["finalStatus"] = v
        elif k == "total_days": res["totalDays"] = v
        elif k == "present_days": res["presentDays"] = v
        elif k == "absent_days": res["absentDays"] = v
        elif k == "attendance_percentage": res["attendancePercentage"] = v
        elif k == "communication_skills": res["communicationSkills"] = v
        elif k == "interpersonal_skills": res["interpersonalSkills"] = v
        elif k == "business_etiquette": res["businessEtiquette"] = v
        elif k == "service_orientation": res["serviceOrientation"] = v
        elif k == "emotional_intelligence_empathy": res["emotionalIntelligenceEmpathy"] = v
        elif k == "accountability_ownership": res["accountabilityOwnership"] = v
        elif k == "presentation_skills": res["presentationSkills"] = v
        elif k == "rank": res["rank"] = v
        elif k == "reevaluation_comments": res["reevaluationComments"] = v
        elif k == "reason_for_absence": res["reasonForAbsence"] = v
        elif k == "pc_name": res["pcName"] = v
        elif k == "emp_id": res["empId"] = v
        elif k == "stream_training": res["streamTraining"] = v
        elif k == "comment_reason": res["commentReason"] = v
        elif k == "project_score1_a1": res["projectScore1_a1"] = v
        elif k == "project_score1_a2": res["projectScore1_a2"] = v
        elif k == "project_score2_a1": res["projectScore2_a1"] = v
        elif k == "project_score2_a2": res["projectScore2_a2"] = v
        elif k == "online_coding_a1": res["onlineCoding_a1"] = v
        elif k == "online_coding_a2": res["onlineCoding_a2"] = v
        # MCQ 1-7
        elif k.startswith("mcq") or k.startswith("coding"):
            res[k] = v
        else:
            res[to_camel(k)] = v
    return res

# Helper: Convert camelCase request dicts to snake_case for DB storage
def map_api_to_db(data: dict) -> dict:
    res = {}
    for k, v in data.items():
        # Specific mappings
        if k == "contactNumber": res["contact_number"] = v
        elif k == "foundationLanguage": res["foundation_language"] = v
        elif k == "emailSentDate": res["email_sent_date"] = v
        elif k == "trainingStatus": res["training_status"] = v
        elif k == "supersetId": res["superset_id"] = v
        elif k == "batchId": res["batch_id"] = v
        elif k == "candidateId": res["candidate_id"] = v
        elif k == "trainerName": res["trainer_name"] = v
        elif k == "trainingStartDate": res["training_start_date"] = v
        elif k == "trainingEndDate": res["training_end_date"] = v
        elif k == "batchNo": res["batch_no"] = v
        elif k == "a1Score": res["a1_score"] = v
        elif k == "a2Score": res["a2_score"] = v
        elif k == "finalStatus": res["final_status"] = v
        elif k == "totalDays": res["total_days"] = v
        elif k == "presentDays": res["present_days"] = v
        elif k == "absentDays": res["absent_days"] = v
        elif k == "attendancePercentage": res["attendance_percentage"] = v
        elif k == "communicationSkills": res["communication_skills"] = v
        elif k == "interpersonalSkills": res["interpersonal_skills"] = v
        elif k == "businessEtiquette": res["business_etiquette"] = v
        elif k == "serviceOrientation": res["service_orientation"] = v
        elif k == "emotionalIntelligenceEmpathy": res["emotional_intelligence_empathy"] = v
        elif k == "accountabilityOwnership": res["accountability_ownership"] = v
        elif k == "presentationSkills": res["presentation_skills"] = v
        elif k == "rank": res["rank"] = v
        elif k == "reevaluationComments": res["reevaluation_comments"] = v
        elif k == "reasonForAbsence": res["reason_for_absence"] = v
        elif k == "pcName": res["pc_name"] = v
        elif k == "empId": res["emp_id"] = v
        elif k == "streamTraining": res["stream_training"] = v
        elif k == "commentReason": res["comment_reason"] = v
        elif k == "projectEval_a1": res["project_eval_a1"] = v
        elif k == "projectEval_a2": res["project_eval_a2"] = v
        elif k == "finalGrade_a1": res["final_grade_a1"] = v
        elif k == "finalGrade_a2": res["final_grade_a2"] = v
        elif k == "projectScore1_a1": res["project_score1_a1"] = v
        elif k == "projectScore1_a2": res["project_score1_a2"] = v
        elif k == "projectScore2_a1": res["project_score2_a1"] = v
        elif k == "projectScore2_a2": res["project_score2_a2"] = v
        elif k == "onlineCoding_a1": res["online_coding_a1"] = v
        elif k == "onlineCoding_a2": res["online_coding_a2"] = v
        # MCQ/Coding 1-7
        elif k.startswith("mcq") or k.startswith("coding"):
            # e.g., mcq1_a1 or mcq1_a2
            # camelCase is mcq1A1 or mcq1A2 -> convert to mcq1_a1/mcq1_a2
            snake = k
            for i in range(1, 8):
                if f"mcq{i}A1" == k: snake = f"mcq{i}_a1"
                elif f"mcq{i}A2" == k: snake = f"mcq{i}_a2"
                elif f"coding{i}A1" == k: snake = f"coding{i}_a1"
                elif f"coding{i}A2" == k: snake = f"coding{i}_a2"
            res[snake] = v
        else:
            # General convert to snake_case
            import re
            snake = re.sub(r'(?<!^)(?=[A-Z])', '_', k).lower()
            res[snake] = v
    return res


def resolve_attempt(a1: Optional[float], a2: Optional[float], pass_mark: float = 60.0) -> Optional[float]:
    if a1 is None:
        return None
    if a1 >= pass_mark:
        return a1
    if a2 is not None:
        return min(a2, pass_mark)
    return a1


# ==========================================
# Spark Phase 1 Endpoints
# ==========================================

@router.get("/spark1/{batch_id}", response_model=List[SparkReportCardResponse])
async def get_spark1_records(batch_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    res = db.table("spark_1_report_cards").select("*").eq("batch_id", batch_id).execute()
    return [SparkReportCardResponse(**map_db_to_api(row)) for row in res.data or []]

@router.put("/spark1/{id}", response_model=SparkReportCardResponse)
async def update_spark1_record(id: str, payload: SparkReportCardUpdate, current_user: dict = Depends(has_role("TRAINER", "COORDINATOR", "ADMIN"))):
    db = get_db()
    check_report_card_access(db, current_user, "spark_1_report_cards", id)
    
    # 1. Fetch current record
    current_res = db.table("spark_1_report_cards").select("*").eq("id", id).execute()
    if not current_res.data:
        raise HTTPException(status_code=404, detail="Record not found")
    current_record = current_res.data[0]
    
    batch_id = current_record.get("batch_id")
    if batch_id:
        batch_res = db.table("batches").select("*").eq("id", batch_id).execute()
        if batch_res.data:
            from app.routers.batch import sync_batch_status
            batch = sync_batch_status(db, batch_res.data[0])
            if batch.get("status") == "CLOSED":
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Cannot update report card for a CLOSED batch."
                )
    
    # 2. Merge with updates
    db_update = map_api_to_db(payload.model_dump(exclude_unset=True))
    merged = {**current_record, **db_update}
    
    # 3. Calculate attendance percentage automatically
    total_d = merged.get("total_days") or 0
    present_d = merged.get("present_days") or 0
    if total_d > 0:
        db_update["attendance_percentage"] = (present_d / total_d) * 100.0
        merged["attendance_percentage"] = db_update["attendance_percentage"]
    
    # 4. Enforce automated clearance and soft skills grading logic
    soft_skills_keys = [
        "communication_skills", "interpersonal_skills", "business_etiquette",
        "service_orientation", "emotional_intelligence_empathy",
        "accountability_ownership", "presentation_skills"
    ]
    soft_skills_vals = [merged.get(k) for k in soft_skills_keys if merged.get(k) is not None]
    
    has_failed = False
    is_in_progress = False
    
    # Attendance Check (80%)
    att_pct = merged.get("attendance_percentage") or 0.0
    if total_d > 0 and att_pct < 80.0:
        has_failed = True
        
    # Technical Check (A1 and A2)
    a1_sc = merged.get("a1_score")
    a2_sc = merged.get("a2_score")
    if a1_sc is not None:
        if a1_sc < 60.0:
            has_failed = True
    else:
        is_in_progress = True
        
    if a2_sc is not None:
        if a2_sc < 60.0:
            has_failed = True
    else:
        is_in_progress = True
        
    # Soft skills Check
    if soft_skills_vals:
        avg_soft = sum(soft_skills_vals) / len(soft_skills_vals)
        if avg_soft < 3.5:
            has_failed = True
    else:
        is_in_progress = True
        
    if has_failed:
        db_update["final_status"] = "Failed"
    elif is_in_progress:
        db_update["final_status"] = "Not Cleared"
    else:
        db_update["final_status"] = "Cleared"
        
    res = db.table("spark_1_report_cards").update(db_update).eq("id", id).execute()
    clean_and_sync_pool(db)
    return SparkReportCardResponse(**map_db_to_api(res.data[0]))

@router.get("/spark1/{batch_id}/download")
async def download_spark1_sheet(batch_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    batch_res = db.table("batches").select("batch_name").eq("id", batch_id).execute()
    batch_name = batch_res.data[0]["batch_name"] if batch_res.data else "Cohort"
    
    records_res = db.table("spark_1_report_cards").select("*").eq("batch_id", batch_id).execute()
    records = [map_db_to_api(r) for r in records_res.data or []]
    records.sort(key=lambda x: x.get("name", "").lower())

    wb = Workbook()
    ws = wb.active
    ws.title = "Spark Phase 1"
    ws.views.sheetView[0].showGridLines = True

    # Styling
    title_font = Font(name="Calibri", size=16, bold=True, color="1F497D")
    bold_font = Font(name="Calibri", size=10, bold=True)
    reg_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    thin_side = Side(border_style="thin", color="D3D3D3")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # Curated fills
    personal_fill = PatternFill(start_color="F5EEF8", end_color="F5EEF8", fill_type="solid")
    rating_fill = PatternFill(start_color="A9DFBF", end_color="A9DFBF", fill_type="solid")    # Light green
    attendance_fill = PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid") # Yellow

    # Title row
    ws.merge_cells("A1:S1")
    ws["A1"] = "Spark Phase 1 Report Card - " + batch_name
    ws["A1"].font = title_font
    ws["A1"].alignment = center_align

    # Row 2: Spanning group headers
    ws.merge_cells("A2:K2")
    ws["A2"] = "Personal & Training Info"
    ws.merge_cells("L2:O2")
    ws["L2"] = "Performance Metrics"
    ws.merge_cells("P2:S2")
    ws["P2"] = "Attendance"

    # Style row 2
    for col in range(1, 20):
        cell = ws.cell(row=2, column=col)
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.alignment = center_align
        cell.border = thin_border
        if col <= 11:
            cell.fill = personal_fill
        elif col <= 15:
            cell.fill = rating_fill
        else:
            cell.fill = attendance_fill

    # Row 3 & 4: Merges and labels
    vertical_merge_cols = [
        (1, "S.No"),
        (2, "Superset ID"),
        (3, "Name"),
        (4, "Registered"),
        (5, "College"),
        (6, "Training"),
        (7, "Training Start date"),
        (8, "Training End date"),
        (9, "Trainer Name"),
        (10, "Batch No"),
        (11, "Training Status"),
        (14, "Final Status"),
        (15, "Comments"),
        (16, "Total"),
        (17, "No."),
        (18, "No."),
        (19, "Atte")
    ]
    for col_idx, val in vertical_merge_cols:
        ws.merge_cells(start_row=3, start_column=col_idx, end_row=4, end_column=col_idx)
        ws.cell(row=3, column=col_idx, value=val)

    # Horizontal merge for A-1/A-2 parent header
    ws.merge_cells("L3:M3")
    ws["L3"] = "Spark Phase 1"

    # Individual sub-headers in Row 4 for scores
    ws.cell(row=4, column=12, value="A-1")
    ws.cell(row=4, column=13, value="A-2")

    # Style row 3 and 4
    for r in [3, 4]:
        for col in range(1, 20):
            cell = ws.cell(row=r, column=col)
            cell.font = bold_font
            cell.alignment = center_align
            cell.border = thin_border
            if col <= 11:
                cell.fill = personal_fill
            elif col <= 15:
                cell.fill = rating_fill
            else:
                cell.fill = attendance_fill

    # Column widths
    column_widths = {
        1: 6,   # S.No
        2: 14,  # Superset ID
        3: 22,  # Name
        4: 25,  # Registered
        5: 22,  # College
        6: 18,  # Training
        7: 18,  # Training Start date
        8: 18,  # Training End date
        9: 18,  # Trainer Name
        10: 12, # Batch No
        11: 16, # Training Status
        12: 10, # A-1
        13: 10, # A-2
        14: 16, # Final Status
        15: 24, # Comments
        16: 10, # Total
        17: 10, # No. (Present)
        18: 10, # No. (Absent)
        19: 12, # Atte (Attendance %)
    }
    for col_idx, width in column_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Data Validation for Training Status dropdown (col K = 11)
    status_dv = DataValidation(type="list", formula1='"Active,Delayed,Discontinued,Not-Cleared,Onboarded,Pending"', allow_blank=True)
    status_dv.error = "Please select a valid Training Status"
    status_dv.errorTitle = "Invalid Status"
    ws.add_data_validation(status_dv)

    # Fill data rows starting at row 5
    for r_idx, r in enumerate(records, 5):
        ws.cell(row=r_idx, column=1, value=r_idx - 4).alignment = center_align
        ws.cell(row=r_idx, column=2, value=r.get("supersetId", "")).alignment = center_align
        ws.cell(row=r_idx, column=3, value=r.get("name", "")).alignment = left_align
        ws.cell(row=r_idx, column=4, value=r.get("email", "")).alignment = left_align
        ws.cell(row=r_idx, column=5, value=r.get("college", "")).alignment = left_align
        ws.cell(row=r_idx, column=6, value=r.get("trainingName", "Spark Phase 1")).alignment = center_align
        ws.cell(row=r_idx, column=7, value=r.get("trainingStartDate", "")).alignment = center_align
        ws.cell(row=r_idx, column=8, value=r.get("trainingEndDate", "")).alignment = center_align
        ws.cell(row=r_idx, column=9, value=r.get("trainerName", "")).alignment = left_align
        ws.cell(row=r_idx, column=10, value=r.get("batchNo", "")).alignment = center_align
        
        status_cell = ws.cell(row=r_idx, column=11, value=r.get("trainingStatus", "Active"))
        status_cell.alignment = center_align
        status_dv.add(status_cell)

        ws.cell(row=r_idx, column=12, value=r.get("a1Score")).alignment = center_align
        ws.cell(row=r_idx, column=13, value=r.get("a2Score")).alignment = center_align
        ws.cell(row=r_idx, column=14, value=r.get("finalStatus", "Not Cleared")).alignment = center_align
        ws.cell(row=r_idx, column=15, value=r.get("reevaluationComments", "")).alignment = left_align
        ws.cell(row=r_idx, column=16, value=r.get("totalDays", 0)).alignment = center_align
        ws.cell(row=r_idx, column=17, value=r.get("presentDays", 0)).alignment = center_align
        ws.cell(row=r_idx, column=18, value=r.get("absentDays", 0)).alignment = center_align
        
        pct_cell = ws.cell(row=r_idx, column=19, value=f"=IF(P{r_idx}>0, Q{r_idx}/P{r_idx}, 0)")
        pct_cell.alignment = center_align
        pct_cell.number_format = '0%'

        for c_idx in range(1, 20):
            ws.cell(row=r_idx, column=c_idx).border = thin_border
            ws.cell(row=r_idx, column=c_idx).font = reg_font

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    
    filename = f"Spark1_Report_{batch_name.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        file_stream,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'}
    )

@router.post("/spark1/{batch_id}/upload")
async def upload_spark1_sheet(batch_id: str, file: UploadFile = File(...), current_user: dict = Depends(has_role("TRAINER", "COORDINATOR", "ADMIN"))):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    
    batch_name = "Unknown"
    batch_res = db.table("batches").select("*").eq("id", batch_id).execute()
    if batch_res.data:
        batch_name = batch_res.data[0].get("batch_name", "Unknown")
        from app.routers.batch import sync_batch_status
        batch = sync_batch_status(db, batch_res.data[0])
        if batch.get("status") == "CLOSED":
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Cannot upload report card for a CLOSED batch."
            )
    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    ws = wb.active
    
    updated_count = 0
    errors = []

    def float_or_none(val):
        try: return float(val) if val is not None else None
        except: return None

    def int_or_zero(val):
        try: return int(val) if val is not None else 0
        except: return 0

    def str_or_none(val):
        return str(val).strip() if val is not None else None

    # Pre-fetch existing records and candidate IDs
    existing_res = db.table("spark_1_report_cards").select("id, email, candidate_id").eq("batch_id", batch_id).execute()
    email_to_id = {r["email"].strip().lower(): r["id"] for r in existing_res.data} if existing_res.data else {}
    email_to_cand_id = {r["email"].strip().lower(): r["candidate_id"] for r in existing_res.data} if existing_res.data else {}

    cand_res = db.table("candidates").select("id, email").eq("batch_id", batch_id).execute()
    db_email_to_cand_id = {c["email"].strip().lower(): c["id"] for c in cand_res.data} if cand_res.data else {}

    payloads = []
    # Iterate rows starting at row 5 (row 3 & 4 = headers)
    for r_idx in range(5, ws.max_row + 1):
        email = ws.cell(row=r_idx, column=4).value  # D = Registered Email
        if not email:
            continue
        email = str(email).strip().lower()
        
        try:
            payload = {
                "superset_id": str_or_none(ws.cell(row=r_idx, column=2).value),          # B
                "name": str_or_none(ws.cell(row=r_idx, column=3).value),                 # C
                "college": str_or_none(ws.cell(row=r_idx, column=5).value),              # E
                "training_name": str_or_none(ws.cell(row=r_idx, column=6).value) or "Spark Phase 1", # F
                "training_start_date": str_or_none(ws.cell(row=r_idx, column=7).value),  # G
                "training_end_date": str_or_none(ws.cell(row=r_idx, column=8).value),    # H
                "trainer_name": str_or_none(ws.cell(row=r_idx, column=9).value),          # I
                "batch_no": str_or_none(ws.cell(row=r_idx, column=10).value),            # J
                "training_status": str_or_none(ws.cell(row=r_idx, column=11).value) or "Active",  # K
                "a1_score": float_or_none(ws.cell(row=r_idx, column=12).value),          # L
                "a2_score": float_or_none(ws.cell(row=r_idx, column=13).value),          # M
                "final_status": str_or_none(ws.cell(row=r_idx, column=14).value) or "Not Cleared", # N
                "reevaluation_comments": str_or_none(ws.cell(row=r_idx, column=15).value), # O
                "total_days": int_or_zero(ws.cell(row=r_idx, column=16).value),            # P
                "present_days": int_or_zero(ws.cell(row=r_idx, column=17).value),          # Q
                "absent_days": int_or_zero(ws.cell(row=r_idx, column=18).value),           # R
                "batch_id": batch_id,
                "email": email
            }
            
            if payload["total_days"] > 0:
                payload["attendance_percentage"] = payload["present_days"] / payload["total_days"]

            candidate_id = email_to_cand_id.get(email) or db_email_to_cand_id.get(email)

            if email in email_to_id:
                payload["id"] = email_to_id[email]
                payloads.append(payload)
            elif candidate_id:
                payload["candidate_id"] = candidate_id
                payloads.append(payload)
            else:
                errors.append(f"Row {r_idx} (email: {email}): Candidate profile not found in this batch")
            
        except Exception as e:
            errors.append(f"Row {r_idx} (email: {email}): {str(e)}")

    # Bulk upsert in chunks of 500
    chunk_size = 500
    for i in range(0, len(payloads), chunk_size):
        chunk = payloads[i:i + chunk_size]
        try:
            db.table("spark_1_report_cards").upsert(chunk).execute()
            updated_count += len(chunk)
        except Exception as e:
            errors.append(f"Bulk upload error (batch {i//chunk_size + 1}): {str(e)}")

    clean_and_sync_pool(db)
    try:
        from app.core.logging_helper import log_file_upload_and_notify
        log_file_upload_and_notify(
            user=current_user,
            filename=file.filename,
            file_type="SPARK1_REPORT_CARD",
            batch_id=batch_id,
            batch_name=batch_name,
            row_count=updated_count,
            status="SUCCESS"
        )
    except Exception as log_err:
        print(f"[Warn] Failed to log success: {log_err}")
    return {"updated": updated_count, "errors": errors}


# ==========================================
# Spark Phase 2 Endpoints
# ==========================================

@router.get("/spark2/{batch_id}", response_model=List[SparkReportCardResponse])
async def get_spark2_records(batch_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    res = db.table("spark_2_report_cards").select("*").eq("batch_id", batch_id).execute()
    return [SparkReportCardResponse(**map_db_to_api(row)) for row in res.data or []]

@router.put("/spark2/{id}", response_model=SparkReportCardResponse)
async def update_spark2_record(id: str, payload: SparkReportCardUpdate, current_user: dict = Depends(has_role("TRAINER", "COORDINATOR", "ADMIN"))):
    db = get_db()
    check_report_card_access(db, current_user, "spark_2_report_cards", id)
    
    # 1. Fetch current record
    current_res = db.table("spark_2_report_cards").select("*").eq("id", id).execute()
    if not current_res.data:
        raise HTTPException(status_code=404, detail="Record not found")
    current_record = current_res.data[0]
    
    batch_id = current_record.get("batch_id")
    if batch_id:
        batch_res = db.table("batches").select("*").eq("id", batch_id).execute()
        if batch_res.data:
            from app.routers.batch import sync_batch_status
            batch = sync_batch_status(db, batch_res.data[0])
            if batch.get("status") == "CLOSED":
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Cannot update report card for a CLOSED batch."
                )
    
    # 2. Merge with updates
    db_update = map_api_to_db(payload.model_dump(exclude_unset=True))
    merged = {**current_record, **db_update}
    
    # 3. Calculate attendance percentage automatically
    total_d = merged.get("total_days") or 0
    present_d = merged.get("present_days") or 0
    if total_d > 0:
        db_update["attendance_percentage"] = (present_d / total_d) * 100.0
        merged["attendance_percentage"] = db_update["attendance_percentage"]
    
    # 4. Enforce automated clearance and soft skills grading logic
    soft_skills_keys = [
        "communication_skills", "interpersonal_skills", "business_etiquette",
        "service_orientation", "emotional_intelligence_empathy",
        "accountability_ownership", "presentation_skills"
    ]
    soft_skills_vals = [merged.get(k) for k in soft_skills_keys if merged.get(k) is not None]
    
    has_failed = False
    is_in_progress = False
    
    # Attendance Check (80%)
    att_pct = merged.get("attendance_percentage") or 0.0
    if total_d > 0 and att_pct < 80.0:
        has_failed = True
        
    # Technical Check (A1 and A2)
    a1_sc = merged.get("a1_score")
    a2_sc = merged.get("a2_score")
    if a1_sc is not None:
        if a1_sc < 60.0:
            has_failed = True
    else:
        is_in_progress = True
        
    if a2_sc is not None:
        if a2_sc < 60.0:
            has_failed = True
    else:
        is_in_progress = True
        
    # Soft skills Check
    if soft_skills_vals:
        avg_soft = sum(soft_skills_vals) / len(soft_skills_vals)
        if avg_soft < 3.5:
            has_failed = True
    else:
        is_in_progress = True
        
    if has_failed:
        db_update["final_status"] = "Failed"
    elif is_in_progress:
        db_update["final_status"] = "Not Cleared"
    else:
        db_update["final_status"] = "Cleared"
        
    res = db.table("spark_2_report_cards").update(db_update).eq("id", id).execute()
    clean_and_sync_pool(db)
    return SparkReportCardResponse(**map_db_to_api(res.data[0]))

@router.get("/spark2/{batch_id}/download")
async def download_spark2_sheet(batch_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    batch_res = db.table("batches").select("batch_name").eq("id", batch_id).execute()
    batch_name = batch_res.data[0]["batch_name"] if batch_res.data else "Cohort"
    
    records_res = db.table("spark_2_report_cards").select("*").eq("batch_id", batch_id).execute()
    records = [map_db_to_api(r) for r in records_res.data or []]
    records.sort(key=lambda x: x.get("name", "").lower())

    wb = Workbook()
    ws = wb.active
    ws.title = "Spark Phase 2"
    ws.views.sheetView[0].showGridLines = True

    # Styling
    title_font = Font(name="Calibri", size=16, bold=True, color="1F497D")
    bold_font = Font(name="Calibri", size=10, bold=True)
    reg_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    thin_side = Side(border_style="thin", color="D3D3D3")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # Curated fills
    personal_fill = PatternFill(start_color="F5EEF8", end_color="F5EEF8", fill_type="solid")
    rating_fill = PatternFill(start_color="A9DFBF", end_color="A9DFBF", fill_type="solid")    # Light green
    attendance_fill = PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid") # Yellow

    # Title row
    ws.merge_cells("A1:S1")
    ws["A1"] = "Spark Phase 2 Report Card - " + batch_name
    ws["A1"].font = title_font
    ws["A1"].alignment = center_align

    # Row 2: Spanning group headers
    ws.merge_cells("A2:K2")
    ws["A2"] = "Personal & Training Info"
    ws.merge_cells("L2:O2")
    ws["L2"] = "Performance Metrics"
    ws.merge_cells("P2:S2")
    ws["P2"] = "Attendance"

    # Style row 2
    for col in range(1, 20):
        cell = ws.cell(row=2, column=col)
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.alignment = center_align
        cell.border = thin_border
        if col <= 11:
            cell.fill = personal_fill
        elif col <= 15:
            cell.fill = rating_fill
        else:
            cell.fill = attendance_fill

    # Row 3 & 4: Merges and labels
    vertical_merge_cols = [
        (1, "S.No"),
        (2, "Superset ID"),
        (3, "Name"),
        (4, "Registered"),
        (5, "College"),
        (6, "Training"),
        (7, "Training Start date"),
        (8, "Training End date"),
        (9, "Trainer Name"),
        (10, "Batch No"),
        (11, "Training Status"),
        (14, "Final Status"),
        (15, "Comments"),
        (16, "Total"),
        (17, "No."),
        (18, "No."),
        (19, "Atte")
    ]
    for col_idx, val in vertical_merge_cols:
        ws.merge_cells(start_row=3, start_column=col_idx, end_row=4, end_column=col_idx)
        ws.cell(row=3, column=col_idx, value=val)

    # Horizontal merge for A-1/A-2 parent header
    ws.merge_cells("L3:M3")
    ws["L3"] = "Spark Phase 2"

    # Individual sub-headers in Row 4 for scores
    ws.cell(row=4, column=12, value="A-1")
    ws.cell(row=4, column=13, value="A-2")

    # Style row 3 and 4
    for r in [3, 4]:
        for col in range(1, 20):
            cell = ws.cell(row=r, column=col)
            cell.font = bold_font
            cell.alignment = center_align
            cell.border = thin_border
            if col <= 11:
                cell.fill = personal_fill
            elif col <= 15:
                cell.fill = rating_fill
            else:
                cell.fill = attendance_fill

    # Column widths
    column_widths = {
        1: 6,   # S.No
        2: 14,  # Superset ID
        3: 22,  # Name
        4: 25,  # Registered
        5: 22,  # College
        6: 18,  # Training
        7: 18,  # Training Start date
        8: 18,  # Training End date
        9: 18,  # Trainer Name
        10: 12, # Batch No
        11: 16, # Training Status
        12: 10, # A-1
        13: 10, # A-2
        14: 16, # Final Status
        15: 24, # Comments
        16: 10, # Total
        17: 10, # No. (Present)
        18: 10, # No. (Absent)
        19: 12, # Atte (Attendance %)
    }
    for col_idx, width in column_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Data Validation for Training Status dropdown (col K = 11)
    status_dv = DataValidation(type="list", formula1='"Active,Delayed,Discontinued,Not-Cleared,Onboarded,Pending"', allow_blank=True)
    status_dv.error = "Please select a valid Training Status"
    status_dv.errorTitle = "Invalid Status"
    ws.add_data_validation(status_dv)

    # Fill data rows starting at row 5
    for r_idx, r in enumerate(records, 5):
        ws.cell(row=r_idx, column=1, value=r_idx - 4).alignment = center_align
        ws.cell(row=r_idx, column=2, value=r.get("supersetId", "")).alignment = center_align
        ws.cell(row=r_idx, column=3, value=r.get("name", "")).alignment = left_align
        ws.cell(row=r_idx, column=4, value=r.get("email", "")).alignment = left_align
        ws.cell(row=r_idx, column=5, value=r.get("college", "")).alignment = left_align
        ws.cell(row=r_idx, column=6, value=r.get("trainingName", "Spark Phase 2")).alignment = center_align
        ws.cell(row=r_idx, column=7, value=r.get("trainingStartDate", "")).alignment = center_align
        ws.cell(row=r_idx, column=8, value=r.get("trainingEndDate", "")).alignment = center_align
        ws.cell(row=r_idx, column=9, value=r.get("trainerName", "")).alignment = left_align
        ws.cell(row=r_idx, column=10, value=r.get("batchNo", "")).alignment = center_align
        
        status_cell = ws.cell(row=r_idx, column=11, value=r.get("trainingStatus", "Active"))
        status_cell.alignment = center_align
        status_dv.add(status_cell)

        ws.cell(row=r_idx, column=12, value=r.get("a1Score")).alignment = center_align
        ws.cell(row=r_idx, column=13, value=r.get("a2Score")).alignment = center_align
        ws.cell(row=r_idx, column=14, value=r.get("finalStatus", "Not Cleared")).alignment = center_align
        ws.cell(row=r_idx, column=15, value=r.get("reevaluationComments", "")).alignment = left_align
        ws.cell(row=r_idx, column=16, value=r.get("totalDays", 0)).alignment = center_align
        ws.cell(row=r_idx, column=17, value=r.get("presentDays", 0)).alignment = center_align
        ws.cell(row=r_idx, column=18, value=r.get("absentDays", 0)).alignment = center_align
        
        pct_cell = ws.cell(row=r_idx, column=19, value=f"=IF(P{r_idx}>0, Q{r_idx}/P{r_idx}, 0)")
        pct_cell.alignment = center_align
        pct_cell.number_format = '0%'

        for c_idx in range(1, 20):
            ws.cell(row=r_idx, column=c_idx).border = thin_border
            ws.cell(row=r_idx, column=c_idx).font = reg_font

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    
    filename = f"Spark2_Report_{batch_name.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        file_stream,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'}
    )

@router.post("/spark2/{batch_id}/upload")
async def upload_spark2_sheet(batch_id: str, file: UploadFile = File(...), current_user: dict = Depends(has_role("TRAINER", "COORDINATOR", "ADMIN"))):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    
    batch_name = "Unknown"
    batch_res = db.table("batches").select("*").eq("id", batch_id).execute()
    if batch_res.data:
        batch_name = batch_res.data[0].get("batch_name", "Unknown")
        from app.routers.batch import sync_batch_status
        batch = sync_batch_status(db, batch_res.data[0])
        if batch.get("status") == "CLOSED":
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Cannot upload report card for a CLOSED batch."
            )
    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    ws = wb.active
    
    updated_count = 0
    errors = []

    def float_or_none(val):
        try: return float(val) if val is not None else None
        except: return None

    def int_or_zero(val):
        try: return int(val) if val is not None else 0
        except: return 0

    def str_or_none(val):
        return str(val).strip() if val is not None else None

    # Pre-fetch existing records and candidate IDs
    existing_res = db.table("spark_2_report_cards").select("id, email, candidate_id").eq("batch_id", batch_id).execute()
    email_to_id = {r["email"].strip().lower(): r["id"] for r in existing_res.data} if existing_res.data else {}
    email_to_cand_id = {r["email"].strip().lower(): r["candidate_id"] for r in existing_res.data} if existing_res.data else {}

    cand_res = db.table("candidates").select("id, email").eq("batch_id", batch_id).execute()
    db_email_to_cand_id = {c["email"].strip().lower(): c["id"] for c in cand_res.data} if cand_res.data else {}

    payloads = []
    # Iterate rows starting at row 5 (row 3 & 4 = headers)
    for r_idx in range(5, ws.max_row + 1):
        email = ws.cell(row=r_idx, column=4).value  # D = Registered Email
        if not email:
            continue
        email = str(email).strip().lower()
        
        try:
            payload = {
                "superset_id": str_or_none(ws.cell(row=r_idx, column=2).value),          # B
                "name": str_or_none(ws.cell(row=r_idx, column=3).value),                 # C
                "college": str_or_none(ws.cell(row=r_idx, column=5).value),              # E
                "training_name": str_or_none(ws.cell(row=r_idx, column=6).value) or "Spark Phase 2", # F
                "training_start_date": str_or_none(ws.cell(row=r_idx, column=7).value),  # G
                "training_end_date": str_or_none(ws.cell(row=r_idx, column=8).value),    # H
                "trainer_name": str_or_none(ws.cell(row=r_idx, column=9).value),          # I
                "batch_no": str_or_none(ws.cell(row=r_idx, column=10).value),            # J
                "training_status": str_or_none(ws.cell(row=r_idx, column=11).value) or "Active",  # K
                "a1_score": float_or_none(ws.cell(row=r_idx, column=12).value),          # L
                "a2_score": float_or_none(ws.cell(row=r_idx, column=13).value),          # M
                "final_status": str_or_none(ws.cell(row=r_idx, column=14).value) or "Not Cleared", # N
                "reevaluation_comments": str_or_none(ws.cell(row=r_idx, column=15).value), # O
                "total_days": int_or_zero(ws.cell(row=r_idx, column=16).value),            # P
                "present_days": int_or_zero(ws.cell(row=r_idx, column=17).value),          # Q
                "absent_days": int_or_zero(ws.cell(row=r_idx, column=18).value),           # R
                "batch_id": batch_id,
                "email": email
            }
            
            if payload["total_days"] > 0:
                payload["attendance_percentage"] = payload["present_days"] / payload["total_days"]

            candidate_id = email_to_cand_id.get(email) or db_email_to_cand_id.get(email)

            if email in email_to_id:
                payload["id"] = email_to_id[email]
                payloads.append(payload)
            elif candidate_id:
                payload["candidate_id"] = candidate_id
                payloads.append(payload)
            else:
                errors.append(f"Row {r_idx} (email: {email}): Candidate profile not found in this batch")
            
        except Exception as e:
            errors.append(f"Row {r_idx} (email: {email}): {str(e)}")

    # Bulk upsert in chunks of 500
    chunk_size = 500
    for i in range(0, len(payloads), chunk_size):
        chunk = payloads[i:i + chunk_size]
        try:
            db.table("spark_2_report_cards").upsert(chunk).execute()
            updated_count += len(chunk)
        except Exception as e:
            errors.append(f"Bulk upload error (batch {i//chunk_size + 1}): {str(e)}")

    clean_and_sync_pool(db)
    try:
        from app.core.logging_helper import log_file_upload_and_notify
        log_file_upload_and_notify(
            user=current_user,
            filename=file.filename,
            file_type="SPARK2_REPORT_CARD",
            batch_id=batch_id,
            batch_name=batch_name,
            row_count=updated_count,
            status="SUCCESS"
        )
    except Exception as log_err:
        print(f"[Warn] Failed to log success: {log_err}")
    return {"updated": updated_count, "errors": errors}


# ==========================================
# Foundational Training Endpoints
# ==========================================

@router.get("/foundation/{batch_id}", response_model=List[FoundationReportCardResponse])
async def get_foundation_records(batch_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    res = db.table("foundation_report_cards").select("*").eq("batch_id", batch_id).execute()
    return [FoundationReportCardResponse(**map_db_to_api(row)) for row in res.data or []]

@router.put("/foundation/{id}", response_model=FoundationReportCardResponse)
async def update_foundation_record(id: str, payload: FoundationReportCardUpdate, current_user: dict = Depends(has_role("TRAINER", "COORDINATOR", "ADMIN"))):
    db = get_db()
    check_report_card_access(db, current_user, "foundation_report_cards", id)
    
    # 1. Fetch current record
    current_res = db.table("foundation_report_cards").select("*").eq("id", id).execute()
    if not current_res.data:
        raise HTTPException(status_code=404, detail="Record not found")
    current_record = current_res.data[0]
    
    batch_id = current_record.get("batch_id")
    if batch_id:
        batch_res = db.table("batches").select("*").eq("id", batch_id).execute()
        if batch_res.data:
            from app.routers.batch import sync_batch_status
            batch = sync_batch_status(db, batch_res.data[0])
            if batch.get("status") == "CLOSED":
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Cannot update report card for a CLOSED batch."
                )
    
    # 2. Merge with updates
    db_update = map_api_to_db(payload.model_dump(exclude_unset=True))
    merged = {**current_record, **db_update}
    
    # 3. Resolve attempts for GAs (pass mark = 60.0) and Project (pass mark = 65.0)
    gas_a1 = [merged.get(f"ga{i}_a1") for i in range(1, 6)]
    gas_resolved = []
    
    has_failed = False
    is_in_progress = False
    
    for i in range(1, 6):
        a1 = merged.get(f"ga{i}_a1")
        a2 = merged.get(f"ga{i}_a2")
        
        # Apply option A: Passing Capped Score
        resolved = resolve_attempt(a1, a2, 60.0)
        gas_resolved.append(resolved)
        
        if a1 is not None:
            if a1 < 60.0:
                if a2 is not None:
                    if a2 < 60.0:
                        has_failed = True
                else:
                    is_in_progress = True
        else:
            is_in_progress = True
            
    proj_a1 = merged.get("project_eval_a1")
    proj_a2 = merged.get("project_eval_a2")
    resolved_proj = resolve_attempt(proj_a1, proj_a2, 65.0)
    
    if proj_a1 is not None:
        if proj_a1 < 65.0:
            if proj_a2 is not None:
                if proj_a2 < 65.0:
                    has_failed = True
            else:
                is_in_progress = True
    else:
        is_in_progress = True
        
    # Calculate final grade Attempt 1
    if all(val is not None for val in gas_a1) and proj_a1 is not None:
        ga_a1_avg = sum(gas_a1) / 5.0
        db_update["final_grade_a1"] = (ga_a1_avg * 0.5) + (proj_a1 * 0.5)
        merged["final_grade_a1"] = db_update["final_grade_a1"]
        
    # Calculate final grade Attempt 2 (or final overall resolved grade)
    if all(val is not None for val in gas_resolved) and resolved_proj is not None:
        ga_resolved_avg = sum(gas_resolved) / 5.0
        db_update["final_grade_a2"] = (ga_resolved_avg * 0.5) + (resolved_proj * 0.5)
        merged["final_grade_a2"] = db_update["final_grade_a2"]

    # Calculate status
    if has_failed:
        db_update["training_status"] = "Failed"
    elif is_in_progress:
        db_update["training_status"] = "Active"
    else:
        db_update["training_status"] = "Cleared"
        
    res = db.table("foundation_report_cards").update(db_update).eq("id", id).execute()
    clean_and_sync_pool(db)
    return FoundationReportCardResponse(**map_db_to_api(res.data[0]))

@router.get("/foundation/{batch_id}/download")
async def download_foundation_sheet(batch_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    batch_res = db.table("batches").select("batch_name").eq("id", batch_id).execute()
    batch_name = batch_res.data[0]["batch_name"] if batch_res.data else "Cohort"
    
    records_res = db.table("foundation_report_cards").select("*").eq("batch_id", batch_id).execute()
    records = [map_db_to_api(r) for r in records_res.data or []]
    records.sort(key=lambda x: x.get("name", "").lower())

    wb = Workbook()
    ws = wb.active
    ws.title = "Foundational Training"
    ws.views.sheetView[0].showGridLines = True

    # Style definitions
    title_font = Font(name="Calibri", size=16, bold=True, color="8A3324")
    sec_font = Font(name="Calibri", size=11, bold=True, color="000000")
    bold_font = Font(name="Calibri", size=10, bold=True)
    reg_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    thin_side = Side(border_style="thin", color="D3D3D3")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # Curated fills
    brown_fill = PatternFill(start_color="EDBB99", end_color="EDBB99", fill_type="solid") # Warm Terracotta/Brown
    blue_fill = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")  # Pale Blue

    ws.merge_cells("A1:X1")
    ws["A1"] = "Foundational Training Report Card - " + batch_name
    ws["A1"].font = title_font
    ws["A1"].alignment = center_align

    # Span Headers
    ws.merge_cells("A3:I3")
    ws["A3"] = "Personal & Info"
    ws["A3"].font = sec_font
    ws["A3"].alignment = center_align
    ws["A3"].fill = brown_fill

    ws.merge_cells("J3:W3")
    ws["J3"] = "Grade Assessments & Scores"
    ws["J3"].font = sec_font
    ws["J3"].alignment = center_align
    ws["J3"].fill = blue_fill

    # Row 4 Detail Headers
    headers = [
        ("Superset ID", 14), ("Candidate Name", 20), ("College", 22), ("Personal Mail ID", 25), ("Contact#", 14),
        ("Foundation Programming Language", 24), ("Status", 12), ("Email Sent Date", 16), ("Reason", 16),
        ("GA 1 A1", 9), ("GA 1 A2", 9), ("GA 2 A1", 9), ("GA 2 A2", 9), ("GA 3 A1", 9), ("GA 3 A2", 9),
        ("GA 4 A1", 9), ("GA 4 A2", 9), ("GA 5 A1", 9), ("GA 5 A2", 9), ("Project A1", 11), ("Project A2", 11),
        ("Final A1", 10), ("Final A2", 10), ("Training Status", 14)
    ]

    for idx, (h_name, width) in enumerate(headers, 1):
        cell = ws.cell(row=4, column=idx, value=h_name)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(idx)].width = width
        if idx <= 9:
            cell.fill = brown_fill
        else:
            cell.fill = blue_fill

    for r_idx, r in enumerate(records, 5):
        ws.cell(row=r_idx, column=1, value=r.get("supersetId", "")).alignment = center_align
        ws.cell(row=r_idx, column=2, value=r.get("name", "")).alignment = left_align
        ws.cell(row=r_idx, column=3, value=r.get("college", "")).alignment = left_align
        ws.cell(row=r_idx, column=4, value=r.get("email", "")).alignment = left_align
        ws.cell(row=r_idx, column=5, value=r.get("contactNumber", "")).alignment = center_align
        ws.cell(row=r_idx, column=6, value=r.get("foundationLanguage", "")).alignment = center_align
        ws.cell(row=r_idx, column=7, value=r.get("status", "Active")).alignment = center_align
        
        sent_date = r.get("emailSentDate")
        ws.cell(row=r_idx, column=8, value=sent_date[:10] if sent_date else "").alignment = center_align
        ws.cell(row=r_idx, column=9, value=r.get("reason", "")).alignment = left_align

        ws.cell(row=r_idx, column=10, value=r.get("ga1_a1")).alignment = center_align
        ws.cell(row=r_idx, column=11, value=r.get("ga1_a2")).alignment = center_align
        ws.cell(row=r_idx, column=12, value=r.get("ga2_a1")).alignment = center_align
        ws.cell(row=r_idx, column=13, value=r.get("ga2_a2")).alignment = center_align
        ws.cell(row=r_idx, column=14, value=r.get("ga3_a1")).alignment = center_align
        ws.cell(row=r_idx, column=15, value=r.get("ga3_a2")).alignment = center_align
        ws.cell(row=r_idx, column=16, value=r.get("ga4_a1")).alignment = center_align
        ws.cell(row=r_idx, column=17, value=r.get("ga4_a2")).alignment = center_align
        ws.cell(row=r_idx, column=18, value=r.get("ga5_a1")).alignment = center_align
        ws.cell(row=r_idx, column=19, value=r.get("ga5_a2")).alignment = center_align
        
        ws.cell(row=r_idx, column=20, value=r.get("projectEval_a1")).alignment = center_align
        ws.cell(row=r_idx, column=21, value=r.get("projectEval_a2")).alignment = center_align
        ws.cell(row=r_idx, column=22, value=r.get("finalGrade_a1")).alignment = center_align
        ws.cell(row=r_idx, column=23, value=r.get("finalGrade_a2")).alignment = center_align
        
        ws.cell(row=r_idx, column=24, value=r.get("trainingStatus", "Active")).alignment = center_align

        for c_idx in range(1, 25):
            ws.cell(row=r_idx, column=c_idx).border = thin_border
            ws.cell(row=r_idx, column=c_idx).font = reg_font

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    
    filename = f"Foundation_Report_{batch_name.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        file_stream,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'}
    )

@router.post("/foundation/{batch_id}/upload")
async def upload_foundation_sheet(batch_id: str, file: UploadFile = File(...), current_user: dict = Depends(has_role("TRAINER", "COORDINATOR", "ADMIN"))):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    
    batch_name = "Unknown"
    batch_res = db.table("batches").select("*").eq("id", batch_id).execute()
    if batch_res.data:
        batch_name = batch_res.data[0].get("batch_name", "Unknown")
        from app.routers.batch import sync_batch_status
        batch = sync_batch_status(db, batch_res.data[0])
        if batch.get("status") == "CLOSED":
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Cannot upload report card for a CLOSED batch."
            )
    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents))
    ws = wb.active
    
    updated_count = 0
    errors = []

    # Pre-fetch existing records and candidate IDs
    existing_res = db.table("foundation_report_cards").select("id, email, candidate_id").eq("batch_id", batch_id).execute()
    email_to_id = {r["email"].strip().lower(): r["id"] for r in existing_res.data} if existing_res.data else {}
    email_to_cand_id = {r["email"].strip().lower(): r["candidate_id"] for r in existing_res.data} if existing_res.data else {}

    cand_res = db.table("candidates").select("id, email").eq("batch_id", batch_id).execute()
    db_email_to_cand_id = {c["email"].strip().lower(): c["id"] for c in cand_res.data} if cand_res.data else {}

    payloads = []
    sync_items = []

    for r_idx in range(5, ws.max_row + 1):
        email = ws.cell(row=r_idx, column=4).value
        if not email:
            continue
        email = str(email).strip().lower()
        
        try:
            superset_id = ws.cell(row=r_idx, column=1).value
            college = ws.cell(row=r_idx, column=3).value
            contact_number = ws.cell(row=r_idx, column=5).value
            foundation_language = ws.cell(row=r_idx, column=6).value
            status = ws.cell(row=r_idx, column=7).value
            reason = ws.cell(row=r_idx, column=9).value

            def float_or_none(val):
                try: return float(val) if val is not None else None
                except: return None

            payload = {
                "superset_id": str(superset_id) if superset_id is not None else None,
                "college": str(college) if college is not None else None,
                "contact_number": str(contact_number) if contact_number is not None else None,
                "foundation_language": str(foundation_language) if foundation_language is not None else None,
                "status": str(status) if status is not None else "Active",
                "reason": str(reason) if reason is not None else None,
                "ga1_a1": float_or_none(ws.cell(row=r_idx, column=10).value),
                "ga1_a2": float_or_none(ws.cell(row=r_idx, column=11).value),
                "ga2_a1": float_or_none(ws.cell(row=r_idx, column=12).value),
                "ga2_a2": float_or_none(ws.cell(row=r_idx, column=13).value),
                "ga3_a1": float_or_none(ws.cell(row=r_idx, column=14).value),
                "ga3_a2": float_or_none(ws.cell(row=r_idx, column=15).value),
                "ga4_a1": float_or_none(ws.cell(row=r_idx, column=16).value),
                "ga4_a2": float_or_none(ws.cell(row=r_idx, column=17).value),
                "ga5_a1": float_or_none(ws.cell(row=r_idx, column=18).value),
                "ga5_a2": float_or_none(ws.cell(row=r_idx, column=19).value),
                "project_eval_a1": float_or_none(ws.cell(row=r_idx, column=20).value),
                "project_eval_a2": float_or_none(ws.cell(row=r_idx, column=21).value),
                "final_grade_a1": float_or_none(ws.cell(row=r_idx, column=22).value),
                "final_grade_a2": float_or_none(ws.cell(row=r_idx, column=23).value),
                "training_status": str(ws.cell(row=r_idx, column=24).value or "Active"),
                "batch_id": batch_id,
                "email": email
            }

            # Handle email sent date parse
            sent_date_raw = ws.cell(row=r_idx, column=8).value
            if sent_date_raw:
                try:
                    payload["email_sent_date"] = datetime.strptime(str(sent_date_raw), "%Y-%m-%d").isoformat()
                except:
                    try:
                        payload["email_sent_date"] = datetime.fromisoformat(str(sent_date_raw).replace('Z', '+00:00')).isoformat()
                    except:
                        pass

            candidate_id = email_to_cand_id.get(email) or db_email_to_cand_id.get(email)

            if email in email_to_id:
                payload["id"] = email_to_id[email]
                payloads.append(payload)
            elif candidate_id:
                payload["candidate_id"] = candidate_id
                payloads.append(payload)
            else:
                errors.append(f"Row {r_idx} (email: {email}): Candidate profile not found in this batch")
                continue

            if candidate_id:
                # Accumulate sync items
                for i in range(1, 6):
                    for att in ["a1", "a2"]:
                        score_key = f"ga{i}_{att}"
                        label = f"GA{i} - Attempt {att[-1]}"
                        if payload.get(score_key) is not None:
                            sync_items.append({
                                "candidate_id": candidate_id,
                                "assessment_name": label,
                                "obtained_score": payload[score_key]
                            })
                if payload.get("project_eval_a1") is not None:
                    sync_items.append({
                        "candidate_id": candidate_id,
                        "assessment_name": "Project Evaluation - Attempt 1",
                        "obtained_score": payload["project_eval_a1"]
                    })
                if payload.get("project_eval_a2") is not None:
                    sync_items.append({
                        "candidate_id": candidate_id,
                        "assessment_name": "Project Evaluation - Attempt 2",
                        "obtained_score": payload["project_eval_a2"]
                    })
                if payload.get("final_grade_a1") is not None:
                    sync_items.append({
                        "candidate_id": candidate_id,
                        "assessment_name": "Final Grade - Attempt 1",
                        "obtained_score": payload["final_grade_a1"]
                    })
                if payload.get("final_grade_a2") is not None:
                    sync_items.append({
                        "candidate_id": candidate_id,
                        "assessment_name": "Final Grade - Attempt 2",
                        "obtained_score": payload["final_grade_a2"]
                    })
        except Exception as e:
            errors.append(f"Row {r_idx} (email: {email}): {str(e)}")

    # Bulk upsert report cards
    chunk_size = 500
    for i in range(0, len(payloads), chunk_size):
        chunk = payloads[i:i + chunk_size]
        try:
            db.table("foundation_report_cards").upsert(chunk).execute()
            updated_count += len(chunk)
        except Exception as e:
            errors.append(f"Bulk upload error (batch {i//chunk_size + 1}): {str(e)}")

    # Bulk sync assessments
    if sync_items:
        try:
            from app.services.assessment_sync_service import AssessmentSyncService
            AssessmentSyncService.bulk_sync_report_card_to_assessments(db, batch_id, sync_items)
        except Exception as sync_err:
            print(f"[Warn] Failed bulk reverse sync for Foundation: {sync_err}")

    clean_and_sync_pool(db)
    try:
        from app.core.logging_helper import log_file_upload_and_notify
        log_file_upload_and_notify(
            user=current_user,
            filename=file.filename,
            file_type="FOUNDATION_REPORT_CARD",
            batch_id=batch_id,
            batch_name=batch_name,
            row_count=updated_count,
            status="SUCCESS"
        )
    except Exception as log_err:
        print(f"[Warn] Failed to log success: {log_err}")
    return {"updated": updated_count, "errors": errors}


# ==========================================
# Stream Based Training Endpoints
# ==========================================

@router.get("/stream/{batch_id}", response_model=List[StreamReportCardResponse])
async def get_stream_records(batch_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    res = db.table("stream_report_cards").select("*").eq("batch_id", batch_id).execute()
    return [StreamReportCardResponse(**map_db_to_api(row)) for row in res.data or []]

@router.put("/stream/{id}", response_model=StreamReportCardResponse)
async def update_stream_record(id: str, payload: StreamReportCardUpdate, current_user: dict = Depends(has_role("TRAINER", "COORDINATOR", "ADMIN"))):
    db = get_db()
    check_report_card_access(db, current_user, "stream_report_cards", id)
    
    # 1. Fetch current record
    current_res = db.table("stream_report_cards").select("*").eq("id", id).execute()
    if not current_res.data:
        raise HTTPException(status_code=404, detail="Record not found")
    current_record = current_res.data[0]
    
    batch_id = current_record.get("batch_id")
    if batch_id:
        batch_res = db.table("batches").select("*").eq("id", batch_id).execute()
        if batch_res.data:
            from app.routers.batch import sync_batch_status
            batch = sync_batch_status(db, batch_res.data[0])
            if batch.get("status") == "CLOSED":
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Cannot update report card for a CLOSED batch."
                )
    
    # 2. Merge with updates
    db_update = map_api_to_db(payload.model_dump(exclude_unset=True))
    merged = {**current_record, **db_update}
    
    # 3. Calculate attendance percentage automatically
    total_d = merged.get("total_days") or 0
    present_d = merged.get("present_days") or 0
    if total_d > 0:
        db_update["attendance_percentage"] = (present_d / total_d) * 100.0
        merged["attendance_percentage"] = db_update["attendance_percentage"]

    # 4. Resolve attempts and check passing criteria
    has_failed = False
    is_in_progress = False
    
    # Attendance Check (85%)
    att_pct = merged.get("attendance_percentage") or 0.0
    if total_d > 0 and att_pct < 85.0:
        has_failed = True
        
    # MCQs Check (MCQ 1-7, pass mark = 60.0)
    for i in range(1, 8):
        a1 = merged.get(f"mcq{i}_a1")
        a2 = merged.get(f"mcq{i}_a2")
        if a1 is not None:
            if a1 < 60.0:
                if a2 is not None:
                    if a2 < 60.0:
                        has_failed = True
                else:
                    is_in_progress = True
        else:
            is_in_progress = True
            
    # Coding Check (Coding 1-7, pass mark = 60.0)
    for i in range(1, 8):
        a1 = merged.get(f"coding{i}_a1")
        a2 = merged.get(f"coding{i}_a2")
        if a1 is not None:
            if a1 < 60.0:
                if a2 is not None:
                    if a2 < 60.0:
                        has_failed = True
                else:
                    is_in_progress = True
        else:
            is_in_progress = True
            
    # Projects Check (Project 1-2, pass mark = 65.0)
    for i in range(1, 3):
        a1 = merged.get(f"project_score{i}_a1")
        a2 = merged.get(f"project_score{i}_a2")
        if a1 is not None:
            if a1 < 65.0:
                if a2 is not None:
                    if a2 < 65.0:
                        has_failed = True
                else:
                    is_in_progress = True
        else:
            is_in_progress = True
            
    # Online Coding Check (pass mark = 60.0)
    oc_a1 = merged.get("online_coding_a1")
    oc_a2 = merged.get("online_coding_a2")
    if oc_a1 is not None:
        if oc_a1 < 60.0:
            if oc_a2 is not None:
                if oc_a2 < 60.0:
                    has_failed = True
            else:
                is_in_progress = True
    else:
        is_in_progress = True

    # 5. Set final status
    if has_failed:
        db_update["final_status"] = "Failed"
    elif is_in_progress:
        db_update["final_status"] = "Not Cleared"
    else:
        db_update["final_status"] = "Cleared"
        
    res = db.table("stream_report_cards").update(db_update).eq("id", id).execute()
    clean_and_sync_pool(db)
    return StreamReportCardResponse(**map_db_to_api(res.data[0]))

@router.get("/stream/{batch_id}/download")
async def download_stream_sheet(batch_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    batch_res = db.table("batches").select("batch_name").eq("id", batch_id).execute()
    batch_name = batch_res.data[0]["batch_name"] if batch_res.data else "Cohort"
    
    records_res = db.table("stream_report_cards").select("*").eq("batch_id", batch_id).execute()
    records = [map_db_to_api(r) for r in records_res.data or []]
    records.sort(key=lambda x: x.get("name", "").lower())

    wb = Workbook()
    ws = wb.active
    ws.title = "Stream Based Training"
    ws.views.sheetView[0].showGridLines = True

    # Styling helper variables
    title_font = Font(name="Calibri", size=16, bold=True, color="2C3E50")
    sec_font = Font(name="Calibri", size=11, bold=True, color="000000")
    bold_font = Font(name="Calibri", size=9, bold=True)
    reg_font = Font(name="Calibri", size=9)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    thin_side = Side(border_style="thin", color="D3D3D3")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # curating gorgeous HSL fits
    info_fill = PatternFill(start_color="F5EEF8", end_color="F5EEF8", fill_type="solid") # Info/Personal
    mcq_fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")  # MCQ metrics
    coding_fill = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid") # Coding metrics
    proj_fill = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid") # Project metrics
    att_fill = PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid") # Attendance

    # Title row
    ws.merge_cells("A1:BB1")
    ws["A1"] = "Stream Based Training Report Card - " + batch_name
    ws["A1"].font = title_font
    ws["A1"].alignment = center_align

    # Row 2: Merged group headers
    ws.merge_cells("A2:N2")
    ws["A2"] = "Personal & Training Info"
    ws.merge_cells("O2:AV2")
    ws["O2"] = "Performance Metrics"
    ws.merge_cells("AY2:BB2")
    ws["AY2"] = "Attendance"

    # Style Row 2 group headers
    for col in range(1, 55): # Columns 1 to 54 (A to BB)
        cell = ws.cell(row=2, column=col)
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.alignment = center_align
        cell.border = thin_border
        if col <= 14:
            cell.fill = info_fill
        elif col <= 48:
            cell.fill = coding_fill
        elif col <= 50:
            pass
        else:
            cell.fill = att_fill

    # Row 3 & 4 vertical merges for Personal Info (A to N)
    p_headers = [
        ("S.No", 6), ("DOJ", 12), ("Superset ID", 12), ("Emp ID", 10), ("Name", 18),
        ("Registered Mail ID", 22), ("College", 18), ("Foundation Language", 16),
        ("Stream Training", 18), ("Training Start date", 15), ("Training End date", 15),
        ("Trainer Name", 16), ("Batch No", 10), ("Training Status", 12)
    ]
    col_idx = 1
    for h_name, w in p_headers:
        ws.merge_cells(start_row=3, start_column=col_idx, end_row=4, end_column=col_idx)
        cell = ws.cell(row=3, column=col_idx, value=h_name)
        cell.fill = info_fill
        ws.column_dimensions[get_column_letter(col_idx)].width = w
        col_idx += 1

    # MCQ 1-7 (Row 3 merged pairs, Row 4 attempts)
    for i in range(1, 8):
        ws.merge_cells(start_row=3, start_column=col_idx, end_row=3, end_column=col_idx + 1)
        ws.cell(row=3, column=col_idx, value=f"MCQ-{i}").fill = mcq_fill
        ws.cell(row=4, column=col_idx, value="A1").fill = mcq_fill
        ws.cell(row=4, column=col_idx + 1, value="A2").fill = mcq_fill
        ws.column_dimensions[get_column_letter(col_idx)].width = 7
        ws.column_dimensions[get_column_letter(col_idx + 1)].width = 7
        col_idx += 2

    # Coding 1-7 (Row 3 merged pairs, Row 4 attempts)
    for i in range(1, 8):
        ws.merge_cells(start_row=3, start_column=col_idx, end_row=3, end_column=col_idx + 1)
        ws.cell(row=3, column=col_idx, value=f"Coding-{i}").fill = coding_fill
        ws.cell(row=4, column=col_idx, value="A1").fill = coding_fill
        ws.cell(row=4, column=col_idx + 1, value="A2").fill = coding_fill
        ws.column_dimensions[get_column_letter(col_idx)].width = 8
        ws.column_dimensions[get_column_letter(col_idx + 1)].width = 8
        col_idx += 2

    # Project Scores (Row 3 merged pairs, Row 4 attempts)
    for i in range(1, 3):
        ws.merge_cells(start_row=3, start_column=col_idx, end_row=3, end_column=col_idx + 1)
        ws.cell(row=3, column=col_idx, value=f"Project Score-{i}").fill = proj_fill
        ws.cell(row=4, column=col_idx, value="A1").fill = proj_fill
        ws.cell(row=4, column=col_idx + 1, value="A2").fill = proj_fill
        ws.column_dimensions[get_column_letter(col_idx)].width = 10
        ws.column_dimensions[get_column_letter(col_idx + 1)].width = 10
        col_idx += 2

    # Online Coding (Row 3 merged pairs, Row 4 attempts)
    ws.merge_cells(start_row=3, start_column=col_idx, end_row=3, end_column=col_idx + 1)
    ws.cell(row=3, column=col_idx, value="Online Coding").fill = proj_fill
    ws.cell(row=4, column=col_idx, value="A1").fill = proj_fill
    ws.cell(row=4, column=col_idx + 1, value="A2").fill = proj_fill
    ws.column_dimensions[get_column_letter(col_idx)].width = 11
    ws.column_dimensions[get_column_letter(col_idx + 1)].width = 11
    col_idx += 2

    # Final Status & Comment / Reason (Vertically merged across rows 2, 3, 4)
    ws.merge_cells(start_row=2, start_column=col_idx, end_row=4, end_column=col_idx)
    ws.cell(row=2, column=col_idx, value="Final Status")
    ws.column_dimensions[get_column_letter(col_idx)].width = 14
    col_idx += 1

    ws.merge_cells(start_row=2, start_column=col_idx, end_row=4, end_column=col_idx)
    ws.cell(row=2, column=col_idx, value="Comment / Reason")
    ws.column_dimensions[get_column_letter(col_idx)].width = 22
    col_idx += 1

    # Attendance (Row 3 & 4 vertical merges)
    a_headers = [
        ("Total No of days", 12), ("Total Present days", 12), ("Absentes", 12), ("Percentage", 12)
    ]
    for h_name, w in a_headers:
        ws.merge_cells(start_row=3, start_column=col_idx, end_row=4, end_column=col_idx)
        cell = ws.cell(row=3, column=col_idx, value=h_name)
        cell.fill = att_fill
        ws.column_dimensions[get_column_letter(col_idx)].width = w
        col_idx += 1

    # Apply fonts/alignments/borders to Row 3 and 4 cells
    for r in [3, 4]:
        for c in range(1, col_idx):
            cell = ws.cell(row=r, column=c)
            cell.font = bold_font
            cell.alignment = center_align
            cell.border = thin_border
            if c in [49, 50]:
                cell.font = Font(name="Calibri", size=9, bold=True)
                cell.alignment = center_align
                cell.border = thin_border

    # Fill data rows starting at row 5
    for r_idx, r in enumerate(records, 5):
        ws.cell(row=r_idx, column=1, value=r_idx - 4).alignment = center_align
        
        doj = r.get("doj")
        ws.cell(row=r_idx, column=2, value=doj[:10] if doj else "").alignment = center_align
        ws.cell(row=r_idx, column=3, value=r.get("supersetId", "")).alignment = center_align
        ws.cell(row=r_idx, column=4, value=r.get("empId", "")).alignment = center_align
        ws.cell(row=r_idx, column=5, value=r.get("name", "")).alignment = left_align
        ws.cell(row=r_idx, column=6, value=r.get("email", "")).alignment = left_align
        ws.cell(row=r_idx, column=7, value=r.get("college", "")).alignment = left_align
        ws.cell(row=r_idx, column=8, value=r.get("foundationLanguage", "")).alignment = center_align
        ws.cell(row=r_idx, column=9, value=r.get("streamTraining", "")).alignment = center_align
        ws.cell(row=r_idx, column=10, value=r.get("trainingStartDate", "")).alignment = center_align
        ws.cell(row=r_idx, column=11, value=r.get("trainingEndDate", "")).alignment = center_align
        ws.cell(row=r_idx, column=12, value=r.get("trainerName", "")).alignment = left_align
        ws.cell(row=r_idx, column=13, value=r.get("batchNo", "")).alignment = center_align
        ws.cell(row=r_idx, column=14, value=r.get("trainingStatus", "Active")).alignment = center_align

        # Scores MCQ 1-7 A-1/A-2
        c_score = 15
        for i in range(1, 8):
            ws.cell(row=r_idx, column=c_score, value=r.get(f"mcq{i}_a1")).alignment = center_align
            ws.cell(row=r_idx, column=c_score + 1, value=r.get(f"mcq{i}_a2")).alignment = center_align
            c_score += 2

        # Coding 1-7 A-1/A-2
        for i in range(1, 8):
            ws.cell(row=r_idx, column=c_score, value=r.get(f"coding{i}_a1")).alignment = center_align
            ws.cell(row=r_idx, column=c_score + 1, value=r.get(f"coding{i}_a2")).alignment = center_align
            c_score += 2

        # Projects
        ws.cell(row=r_idx, column=c_score, value=r.get("projectScore1_a1")).alignment = center_align
        ws.cell(row=r_idx, column=c_score + 1, value=r.get("projectScore1_a2")).alignment = center_align
        ws.cell(row=r_idx, column=c_score + 2, value=r.get("projectScore2_a1")).alignment = center_align
        ws.cell(row=r_idx, column=c_score + 3, value=r.get("projectScore2_a2")).alignment = center_align
        c_score += 4

        # Online Coding
        ws.cell(row=r_idx, column=c_score, value=r.get("onlineCoding_a1")).alignment = center_align
        ws.cell(row=r_idx, column=c_score + 1, value=r.get("onlineCoding_a2")).alignment = center_align
        c_score += 2

        # Final Status & Comments
        ws.cell(row=r_idx, column=c_score, value=r.get("finalStatus", "Cleared")).alignment = center_align
        c_score += 1
        ws.cell(row=r_idx, column=c_score, value=r.get("commentReason", "")).alignment = left_align
        c_score += 1

        # Attendance
        ws.cell(row=r_idx, column=c_score, value=r.get("totalDays", 0)).alignment = center_align
        ws.cell(row=r_idx, column=c_score + 1, value=r.get("presentDays", 0)).alignment = center_align
        ws.cell(row=r_idx, column=c_score + 2, value=r.get("absentDays", 0)).alignment = center_align
        
        pct_cell = ws.cell(row=r_idx, column=c_score + 3, value=f"=IF({get_column_letter(c_score)}{r_idx}>0, {get_column_letter(c_score+1)}{r_idx}/{get_column_letter(c_score)}{r_idx}, 0)")
        pct_cell.alignment = center_align
        pct_cell.number_format = '0%'

        for c_idx in range(1, col_idx):
            ws.cell(row=r_idx, column=c_idx).border = thin_border
            ws.cell(row=r_idx, column=c_idx).font = reg_font

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    
    filename = f"Stream_Report_{batch_name.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        file_stream,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'}
    )

@router.post("/stream/{batch_id}/upload")
async def upload_stream_sheet(batch_id: str, file: UploadFile = File(...), current_user: dict = Depends(has_role("TRAINER", "COORDINATOR", "ADMIN"))):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    
    batch_name = "Unknown"
    batch_res = db.table("batches").select("*").eq("id", batch_id).execute()
    if batch_res.data:
        batch_name = batch_res.data[0].get("batch_name", "Unknown")
        from app.routers.batch import sync_batch_status
        batch = sync_batch_status(db, batch_res.data[0])
        if batch.get("status") == "CLOSED":
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Cannot upload report card for a CLOSED batch."
            )
    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents))
    ws = wb.active
    
    updated_count = 0
    errors = []

    # Pre-fetch existing records and candidate IDs
    existing_res = db.table("stream_report_cards").select("id, email, candidate_id").eq("batch_id", batch_id).execute()
    email_to_id = {r["email"].strip().lower(): r["id"] for r in existing_res.data} if existing_res.data else {}
    email_to_cand_id = {r["email"].strip().lower(): r["candidate_id"] for r in existing_res.data} if existing_res.data else {}

    cand_res = db.table("candidates").select("id, email").eq("batch_id", batch_id).execute()
    db_email_to_cand_id = {c["email"].strip().lower(): c["id"] for c in cand_res.data} if cand_res.data else {}

    payloads = []
    sync_items = []

    # Columns index references (derived from layout above)
    # col 5: Name, col 6: Registered Mail ID
    for r_idx in range(5, ws.max_row + 1):
        email = ws.cell(row=r_idx, column=6).value
        if not email:
            continue
        email = str(email).strip().lower()
        
        try:
            superset_id = ws.cell(row=r_idx, column=3).value
            emp_id = ws.cell(row=r_idx, column=4).value
            college = ws.cell(row=r_idx, column=7).value
            foundation_language = ws.cell(row=r_idx, column=8).value
            stream_training = ws.cell(row=r_idx, column=9).value
            training_status = ws.cell(row=r_idx, column=14).value

            def float_or_none(val):
                try: return float(val) if val is not None else None
                except: return None

            def int_or_zero(val):
                try: return int(val) if val is not None else 0
                except: return 0

            # Score extraction starts at column 15
            payload = {
                "superset_id": str(superset_id) if superset_id is not None else None,
                "emp_id": str(emp_id) if emp_id is not None else None,
                "college": str(college) if college is not None else None,
                "foundation_language": str(foundation_language) if foundation_language is not None else None,
                "stream_training": str(stream_training) if stream_training is not None else None,
                "training_status": str(training_status) if training_status is not None else "Active",
                "batch_id": batch_id,
                "email": email
            }

            # MCQ 1-7 A-1/A-2
            c_score = 15
            for i in range(1, 8):
                payload[f"mcq{i}_a1"] = float_or_none(ws.cell(row=r_idx, column=c_score).value)
                payload[f"mcq{i}_a2"] = float_or_none(ws.cell(row=r_idx, column=c_score + 1).value)
                c_score += 2

            # Coding 1-7 A-1/A-2
            for i in range(1, 8):
                payload[f"coding{i}_a1"] = float_or_none(ws.cell(row=r_idx, column=c_score).value)
                payload[f"coding{i}_a2"] = float_or_none(ws.cell(row=r_idx, column=c_score + 1).value)
                c_score += 2

            # Projects
            payload["project_score1_a1"] = float_or_none(ws.cell(row=r_idx, column=c_score).value)
            payload["project_score1_a2"] = float_or_none(ws.cell(row=r_idx, column=c_score + 1).value)
            payload["project_score2_a1"] = float_or_none(ws.cell(row=r_idx, column=c_score + 2).value)
            payload["project_score2_a2"] = float_or_none(ws.cell(row=r_idx, column=c_score + 3).value)
            c_score += 4

            # Online Coding
            payload["online_coding_a1"] = float_or_none(ws.cell(row=r_idx, column=c_score).value)
            payload["online_coding_a2"] = float_or_none(ws.cell(row=r_idx, column=c_score + 1).value)
            c_score += 2

            # Final Status & Comments
            payload["final_status"] = str(ws.cell(row=r_idx, column=c_score).value or "Cleared")
            c_score += 1
            payload["comment_reason"] = str(ws.cell(row=r_idx, column=c_score).value or "")
            c_score += 1

            # Attendance
            payload["total_days"] = int_or_zero(ws.cell(row=r_idx, column=c_score).value)
            payload["present_days"] = int_or_zero(ws.cell(row=r_idx, column=c_score + 1).value)
            payload["absent_days"] = int_or_zero(ws.cell(row=r_idx, column=c_score + 2).value)

            if payload["total_days"] > 0:
                payload["attendance_percentage"] = payload["present_days"] / payload["total_days"]

            # Handle DOJ parse
            doj_raw = ws.cell(row=r_idx, column=2).value
            if doj_raw:
                try:
                    payload["doj"] = datetime.strptime(str(doj_raw), "%Y-%m-%d").isoformat()
                except:
                    try:
                        payload["doj"] = datetime.fromisoformat(str(doj_raw).replace('Z', '+00:00')).isoformat()
                    except:
                        pass

            candidate_id = email_to_cand_id.get(email) or db_email_to_cand_id.get(email)

            if email in email_to_id:
                payload["id"] = email_to_id[email]
                payloads.append(payload)
            elif candidate_id:
                payload["candidate_id"] = candidate_id
                payloads.append(payload)
            else:
                errors.append(f"Row {r_idx} (email: {email}): Candidate profile not found in this batch")
                continue

            if candidate_id:
                # Accumulate sync items
                for i in range(1, 8):
                    for att in ["a1", "a2"]:
                        mcq_key = f"mcq{i}_{att}"
                        coding_key = f"coding{i}_{att}"
                        if payload.get(mcq_key) is not None:
                            sync_items.append({
                                "candidate_id": candidate_id,
                                "assessment_name": f"MCQ {i} - Attempt {att[-1]}",
                                "obtained_score": payload[mcq_key]
                            })
                        if payload.get(coding_key) is not None:
                            sync_items.append({
                                "candidate_id": candidate_id,
                                "assessment_name": f"Coding {i} - Attempt {att[-1]}",
                                "obtained_score": payload[coding_key]
                            })
                for i in range(1, 3):
                    for att in ["a1", "a2"]:
                        proj_key = f"project_score{i}_{att}"
                        if payload.get(proj_key) is not None:
                            sync_items.append({
                                "candidate_id": candidate_id,
                                "assessment_name": f"Project {i} - Attempt {att[-1]}",
                                "obtained_score": payload[proj_key]
                            })
                if payload.get("online_coding_a1") is not None:
                    sync_items.append({
                        "candidate_id": candidate_id,
                        "assessment_name": "Online Coding - Attempt 1",
                        "obtained_score": payload["online_coding_a1"]
                    })
                if payload.get("online_coding_a2") is not None:
                    sync_items.append({
                        "candidate_id": candidate_id,
                        "assessment_name": "Online Coding - Attempt 2",
                        "obtained_score": payload["online_coding_a2"]
                    })
        except Exception as e:
            errors.append(f"Row {r_idx} (email: {email}): {str(e)}")

    # Bulk upsert report cards
    chunk_size = 500
    for i in range(0, len(payloads), chunk_size):
        chunk = payloads[i:i + chunk_size]
        try:
            db.table("stream_report_cards").upsert(chunk).execute()
            updated_count += len(chunk)
        except Exception as e:
            errors.append(f"Bulk upload error (batch {i//chunk_size + 1}): {str(e)}")

    # Bulk sync assessments
    if sync_items:
        try:
            from app.services.assessment_sync_service import AssessmentSyncService
            AssessmentSyncService.bulk_sync_report_card_to_assessments(db, batch_id, sync_items)
        except Exception as sync_err:
            print(f"[Warn] Failed bulk reverse sync for Stream: {sync_err}")

    clean_and_sync_pool(db)
    try:
        from app.core.logging_helper import log_file_upload_and_notify
        log_file_upload_and_notify(
            user=current_user,
            filename=file.filename,
            file_type="STREAM_REPORT_CARD",
            batch_id=batch_id,
            batch_name=batch_name,
            row_count=updated_count,
            status="SUCCESS"
        )
    except Exception as log_err:
        print(f"[Warn] Failed to log success: {log_err}")
    return {"updated": updated_count, "errors": errors}
