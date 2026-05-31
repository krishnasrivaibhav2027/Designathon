from fastapi import APIRouter, HTTPException, status, Depends
from fastapi.responses import StreamingResponse
from typing import List, Optional
from datetime import datetime, timedelta
from app.schemas.schemas import (
    FeedbackCreate, FeedbackUpdate, FeedbackResponse, ToppersListResponse,
    DetailedFeedbackCreate, DetailedFeedbackResponse, FeedbackWindowStatus
)
from app.core.database import get_db
from app.core.security import get_current_user, has_role, check_batch_access, check_candidate_access
from app.models.models import Feedback, row_to_api
from app.services.topper_service import TopperService

router = APIRouter(prefix="/api/report", tags=["report"])


# ─── Helpers ────────────────────────────────────────────────────────────────

FEEDBACK_WINDOW_DAYS_BEFORE = 3   # window opens N days before end date

def _get_feedback_window(end_date_str: str):
    """Return (window_open: bool, opens_on, closes_on) for a batch end date."""
    try:
        end_dt = datetime.fromisoformat(end_date_str.replace("Z", "+00:00")).replace(tzinfo=None)
    except Exception:
        return False, None, None
    opens_on = end_dt - timedelta(days=FEEDBACK_WINDOW_DAYS_BEFORE)
    closes_on = end_dt
    now = datetime.utcnow()
    window_open = opens_on <= now <= closes_on
    return window_open, opens_on, closes_on


def _coordinator_owns_batch(db, batch_id: str, user_id: str) -> bool:
    """Return True if the coordinator created this batch (or is a legacy admin)."""
    batch_res = db.table("batches").select("description").eq("id", batch_id).execute()
    if not batch_res.data:
        return False
    desc_str = batch_res.data[0].get("description")
    creator = ""
    if desc_str and desc_str.startswith("{"):
        try:
            import json
            creator = json.loads(desc_str).get("created_by", "")
        except Exception:
            pass
    legacy_admins = {"df772f20-b396-4a3b-8ddc-68fcd54b6060", "728f45b3-f6bd-4cfa-860f-a42c89682b33"}
    return creator == user_id or (not creator and user_id in legacy_admins)


@router.post("/feedback", response_model=FeedbackResponse)
async def submit_feedback(
    feedback_data: FeedbackCreate,
    current_user: dict = Depends(get_current_user)
):
    """Submit feedback/rating"""
    db = get_db()
    check_batch_access(db, current_user, feedback_data.batchId)
    if current_user.get("role") == "TRAINEE":
        check_candidate_access(db, current_user, feedback_data.candidateId)
    
    try:
        feedback = Feedback(
            batchId=feedback_data.batchId,
            candidateId=feedback_data.candidateId,
            rating=feedback_data.rating,
            comments=feedback_data.comments
        )
        
        result = db.table("feedbacks").insert(feedback.to_dict()).execute()
        
        if not result.data:
            raise HTTPException(status_code=500, detail="Failed to submit feedback")
        
        return FeedbackResponse(**row_to_api(result.data[0]))
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@router.get("/feedback/batch/{batch_id}", response_model=List[FeedbackResponse])
async def get_batch_feedback(
    batch_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Get feedback for batch"""
    db = get_db()
    check_batch_access(db, current_user, batch_id)
                
    try:
        result = db.table("feedbacks").select("*").eq("batch_id", batch_id).execute()
        return [FeedbackResponse(**row_to_api(f)) for f in result.data]
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@router.get("/feedback/candidate/{candidate_id}", response_model=List[FeedbackResponse])
async def get_candidate_feedback(
    candidate_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Get feedback for candidate"""
    db = get_db()
    check_candidate_access(db, current_user, candidate_id)
    
    try:
        result = db.table("feedbacks").select("*").eq("candidate_id", candidate_id).execute()
        return [FeedbackResponse(**row_to_api(f)) for f in result.data]
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@router.get("/toppers/export")
async def export_toppers_list(
    batch_ids: str,
    limit: int = 10,
    current_user: dict = Depends(get_current_user)
):
    db = get_db()
    role = current_user.get("role")
    if role not in ["TRAINER", "COORDINATOR", "ADMIN"]:
        raise HTTPException(status_code=403, detail="Not authorized to access reports")
        
    ids_list = [bid.strip() for bid in batch_ids.split(",") if bid.strip()]
    for bid in ids_list:
        check_batch_access(db, current_user, bid)
        
    try:
        import io
        import openpyxl
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
        
        # 1. Fetch batches
        batches_res = db.table("batches").select("*").in_("id", ids_list).execute()
        batches_map = {b["id"]: b for b in (batches_res.data or [])}
        
        wb = Workbook()
        default_sheet = wb.active
        
        title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=11)
        bold_font = Font(name="Calibri", size=11, bold=True)
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center")
        thin_side = Side(border_style="thin", color="D3D3D3")
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        header_fill = PatternFill(start_color="D4AC0D", end_color="D4AC0D", fill_type="solid")
        alt_fill = PatternFill(start_color="FEF9E7", end_color="FEF9E7", fill_type="solid")
        
        for bid in ids_list:
            batch_name = batches_map.get(bid, {}).get("batch_name", bid)
            toppers = await TopperService.calculate_batch_toppers(bid)
            toppers = toppers[:limit]
            
            sheet_title = batch_name[:30].replace(":", "").replace("/", "").replace("\\", "").replace("?", "").replace("*", "")
            ws = wb.create_sheet(title=sheet_title)
            ws.views.sheetView[0].showGridLines = True
            
            ws.merge_cells("A1:G2")
            title_cell = ws["A1"]
            title_cell.value = f"Toppers Leaderboard - {batch_name}"
            title_cell.font = title_font
            title_cell.alignment = center_align
            title_cell.fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            
            headers = [
                ("Rank", 8),
                ("Registration Number", 18),
                ("Name", 24),
                ("Email", 26),
                ("Overall Score (Option A)", 20),
                ("Assessment Score Avg", 20),
                ("Attendance %", 14)
            ]
            
            for col_idx, (text, width) in enumerate(headers, 1):
                cell = ws.cell(row=4, column=col_idx, value=text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = border
                ws.column_dimensions[get_column_letter(col_idx)].width = width
                
            ws.row_dimensions[4].height = 28
            
            for r_idx, t in enumerate(toppers, 5):
                fill = alt_fill if r_idx % 2 == 0 else None
                row_data = [
                    r_idx - 4,
                    t.get("registrationNumber", "N/A"),
                    t.get("fullName", ""),
                    t.get("email", ""),
                    t.get("overallScore") / 100.0 if t.get("overallScore") is not None else 0.0,
                    t.get("assessmentScore") / 100.0 if t.get("assessmentScore") is not None else 0.0,
                    t.get("attendancePercentage") / 100.0 if t.get("attendancePercentage") is not None else 0.0
                ]
                
                for col_idx, val in enumerate(row_data, 1):
                    cell = ws.cell(row=r_idx, column=col_idx, value=val)
                    cell.font = regular_font if col_idx not in [1, 5] else bold_font
                    cell.border = border
                    if fill:
                        cell.fill = fill
                        
                    if col_idx in [1, 2]:
                        cell.alignment = center_align
                    elif col_idx in [5, 6, 7]:
                        cell.alignment = center_align
                        cell.number_format = '0.0%'
                    else:
                        cell.alignment = left_align
                ws.row_dimensions[r_idx].height = 20
                
        if default_sheet in wb.worksheets and len(wb.worksheets) > 1:
            wb.remove(default_sheet)
            
        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        
        return StreamingResponse(
            file_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="Batch_Toppers_Leaderboard.xlsx"'}
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=400, detail=str(e))


@router.get("/toppers/{batch_id}", response_model=ToppersListResponse)
async def get_batch_toppers(
    batch_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Get toppers for batch"""
    db = get_db()
    check_batch_access(db, current_user, batch_id)
                
    try:
        batch_result = db.table("batches").select("*").eq("id", batch_id).execute()
        if not batch_result.data:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Batch not found"
            )
        
        batch = batch_result.data[0]
        toppers = await TopperService.get_top_performers(batch_id, limit=10)
        
        return ToppersListResponse(
            batchId=batch_id,
            batchName=batch.get("batch_name"),
            toppers=toppers
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@router.get("/rank/candidate/{candidate_id}/batch/{batch_id}")
async def get_candidate_rank(
    candidate_id: str,
    batch_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Get candidate rank in batch"""
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    check_candidate_access(db, current_user, candidate_id)
    try:
        rank_data = await TopperService.get_candidate_rank(batch_id, candidate_id)
        return rank_data
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@router.post("/feedback/request/{batch_id}")
async def request_feedback(
    batch_id: str,
    current_user: dict = Depends(has_role("ADMIN", "COORDINATOR"))
):
    """Trigger sending feedback request emails to all candidates in the batch.
    Respects the feedback window: opens 3 days before end date, closes on end date."""
    db = get_db()
    from app.services.email_service import EmailService
    check_batch_access(db, current_user, batch_id)

    try:
        batch_result = db.table("batches").select("*").eq("id", batch_id).execute()
        if not batch_result.data:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Batch not found")

        batch = batch_result.data[0]
        batch_name = batch.get("batch_name", "")
        end_date_str = batch.get("end_date", "")

        # Check feedback window
        window_open, opens_on, closes_on = _get_feedback_window(end_date_str)
        if not window_open:
            now = datetime.utcnow()
            if opens_on and now < opens_on:
                raise HTTPException(
                    status_code=400,
                    detail=f"Feedback window has not opened yet. It opens on {opens_on.strftime('%d %b %Y')}."
                )
            else:
                raise HTTPException(
                    status_code=400,
                    detail="Feedback window has closed (batch end date has passed)."
                )

        candidates_result = db.table("candidates").select("*").eq("batch_id", batch_id).execute()
        candidates = candidates_result.data or []

        sent_count = 0
        failed_count = 0
        form_url = f"http://localhost:5173/feedback/form?batchId={batch_id}"

        for c in candidates:
            email = c.get("email")
            name = c.get("full_name")
            if email and name:
                success = await EmailService.send_feedback_request(
                    email, name, batch_name,
                    batch_id=batch_id,
                    feedback_form_url=form_url
                )
                if success:
                    sent_count += 1
                else:
                    failed_count += 1

        return {
            "message": "Feedback request process completed",
            "total_candidates": len(candidates),
            "sent": sent_count,
            "failed": failed_count,
            "form_url": form_url,
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@router.get("/progress-tracker")
async def get_progress_tracker(current_user: dict = Depends(get_current_user)):
    """Calculate average trainee assessment scores grouped by course week"""
    db = get_db()
    try:
        # Fetch all batches to get their start_dates
        batches_res = db.table("batches").select("*").execute()
        batches = batches_res.data or []
        
        role = current_user.get("role")
        user_id = current_user.get("sub") or current_user.get("email") or ""
        
        trainer_name = ""
        if role in ("ADMIN", "COORDINATOR", "TRAINER"):
            try:
                user_res = db.table("users").select("full_name").eq("id", current_user.get("sub")).execute()
                if user_res.data:
                    trainer_name = user_res.data[0]["full_name"]
            except Exception:
                pass
        trainer_clean = trainer_name.strip().lower() if trainer_name else ""
        user_email_clean = current_user.get("email", "").strip().lower()
        
        filtered_batches = []
        for b in batches:
            creator = ""
            desc_str = b.get("description")
            if desc_str and desc_str.startswith("{"):
                try:
                    import json
                    creator = json.loads(desc_str).get("created_by", "")
                except:
                    pass
            
            is_allowed = False
            if role in ("ADMIN", "COORDINATOR"):
                legacy_admins = {"df772f20-b396-4a3b-8ddc-68fcd54b6060", "728f45b3-f6bd-4cfa-860f-a42c89682b33"}
                is_owner = creator == user_id or (not creator and user_id in legacy_admins)
                
                trainers = b.get("trainers", []) or []
                is_trainer = False
                for t in trainers:
                    t_clean = t.strip().lower()
                    if (trainer_clean and t_clean == trainer_clean) or t_clean == user_email_clean:
                        is_trainer = True
                        break
                if is_owner or is_trainer:
                    is_allowed = True
                    
            elif role == "TRAINER":
                trainers = b.get("trainers", []) or []
                for t in trainers:
                    t_clean = t.strip().lower()
                    if (trainer_clean and t_clean == trainer_clean) or t_clean == user_email_clean:
                        is_allowed = True
                        break
                        
            elif role == "TRAINEE":
                try:
                    user_res = db.table("users").select("assigned_batches").eq("id", current_user.get("sub")).execute()
                    assigned_batches = user_res.data[0].get("assigned_batches", []) or [] if user_res.data else []
                    if b["id"] in assigned_batches:
                        is_allowed = True
                except:
                    pass
            
            if is_allowed:
                filtered_batches.append(b)
        batches = filtered_batches
            
        batch_map = {b["id"]: b for b in batches}
        batch_ids = list(batch_map.keys())
        if not batch_ids:
            return []
        
        # Fetch only assessments for authorized batches
        assessments_res = db.table("assessments").select("*").in_("batch_id", batch_ids).execute()
        assessments = assessments_res.data or []
        
        # Group assessment scores by week number
        weeks_data = {}
        
        for a in assessments:
            batch_id = a.get("batch_id")
            if not batch_id or batch_id not in batch_map:
                continue
            
            batch = batch_map[batch_id]
            start_date_str = batch.get("start_date")
            created_at_str = a.get("created_at")
            
            if not start_date_str or not created_at_str:
                continue
                
            try:
                # Parse datetimes (handle both Z suffix and other ISO formats)
                start_date = datetime.fromisoformat(start_date_str.replace("Z", "+00:00"))
                created_at = datetime.fromisoformat(created_at_str.replace("Z", "+00:00"))
                
                diff = created_at - start_date
                week_num = max(1, (diff.days // 7) + 1)
                
                if week_num not in weeks_data:
                    weeks_data[week_num] = []
                weeks_data[week_num].append(a.get("percentage", 0))
            except Exception as parse_err:
                print(f"[Warn] Error parsing dates: {parse_err}")
                continue
                
        # Format the result for Recharts
        chart_data = []
        for week_num in sorted(weeks_data.keys()):
            scores = weeks_data[week_num]
            avg_score = sum(scores) / len(scores) if scores else 0
            chart_data.append({
                "week": f"Week {week_num}",
                "score": round(avg_score, 1)
            })
            
        return chart_data
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@router.get("/toppers/all/leaderboard")
async def get_all_toppers(
    limit: int = 5,
    current_user: dict = Depends(get_current_user)
):
    """Get top performing trainees across all batches"""
    db = get_db()
    try:
        # Get all batches
        batches_res = db.table("batches").select("id, batch_name, description, trainers").execute()
        batches = batches_res.data or []
        
        role = current_user.get("role")
        user_id = current_user.get("sub") or current_user.get("email") or ""
        
        trainer_name = ""
        if role in ("ADMIN", "COORDINATOR", "TRAINER"):
            try:
                user_res = db.table("users").select("full_name").eq("id", current_user.get("sub")).execute()
                if user_res.data:
                    trainer_name = user_res.data[0]["full_name"]
            except Exception:
                pass
        trainer_clean = trainer_name.strip().lower() if trainer_name else ""
        user_email_clean = current_user.get("email", "").strip().lower()
        
        filtered_batches = []
        for b in batches:
            creator = ""
            desc_str = b.get("description")
            if desc_str and desc_str.startswith("{"):
                try:
                    import json
                    creator = json.loads(desc_str).get("created_by", "")
                except:
                    pass
            
            is_allowed = False
            if role in ("ADMIN", "COORDINATOR"):
                legacy_admins = {"df772f20-b396-4a3b-8ddc-68fcd54b6060", "728f45b3-f6bd-4cfa-860f-a42c89682b33"}
                is_owner = creator == user_id or (not creator and user_id in legacy_admins)
                
                trainers = b.get("trainers", []) or []
                is_trainer = False
                for t in trainers:
                    t_clean = t.strip().lower()
                    if (trainer_clean and t_clean == trainer_clean) or t_clean == user_email_clean:
                        is_trainer = True
                        break
                if is_owner or is_trainer:
                    is_allowed = True
                    
            elif role == "TRAINER":
                trainers = b.get("trainers", []) or []
                for t in trainers:
                    t_clean = t.strip().lower()
                    if (trainer_clean and t_clean == trainer_clean) or t_clean == user_email_clean:
                        is_allowed = True
                        break
                        
            elif role == "TRAINEE":
                try:
                    user_res = db.table("users").select("assigned_batches").eq("id", current_user.get("sub")).execute()
                    assigned_batches = user_res.data[0].get("assigned_batches", []) or [] if user_res.data else []
                    if b["id"] in assigned_batches:
                        is_allowed = True
                except:
                    pass
            
            if is_allowed:
                filtered_batches.append(b)
        batches = filtered_batches
        
        import asyncio
        all_toppers = []
        
        # Calculate toppers for all batches concurrently
        toppers_results = await asyncio.gather(*[
            TopperService.calculate_batch_toppers(b["id"])
            for b in batches
        ])
        
        for batch, toppers in zip(batches, toppers_results):
            batch_name = batch["batch_name"]
            for t in toppers:
                t["batchName"] = batch_name
                all_toppers.append(t)
                
        # Sort all toppers by overallScore descending
        all_toppers.sort(key=lambda x: x["overallScore"], reverse=True)
        
        # Filter: average attendance > 60% and overall score > 50%
        filtered_toppers = [
            t for t in all_toppers
            if t.get("attendancePercentage", 0) > 60 and t.get("overallScore", 0) > 50
        ]
        
        return filtered_toppers[:limit]
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


# ─── Feedback Window Status ──────────────────────────────────────────────────

@router.get("/feedback/window/{batch_id}", response_model=FeedbackWindowStatus)
async def get_feedback_window(
    batch_id: str,
    current_user: dict = Depends(get_current_user)
):
    """Return whether the feedback window is currently open for a batch."""
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    batch_result = db.table("batches").select("id, batch_name, end_date").eq("id", batch_id).execute()
    if not batch_result.data:
        raise HTTPException(status_code=404, detail="Batch not found")

    batch = batch_result.data[0]
    end_date_str = batch.get("end_date", "")
    window_open, opens_on, closes_on = _get_feedback_window(end_date_str)

    days_until_close = None
    if closes_on:
        delta = closes_on - datetime.utcnow()
        days_until_close = max(0, delta.days)

    return FeedbackWindowStatus(
        batchId=batch_id,
        batchName=batch.get("batch_name", ""),
        endDate=datetime.fromisoformat(end_date_str.replace("Z", "+00:00")),
        windowOpen=window_open,
        windowOpensOn=opens_on,
        windowClosesOn=closes_on,
        daysUntilClose=days_until_close,
    )


# ─── Detailed Feedback Submission (from the feedback form) ──────────────────

@router.post("/feedback/detailed", response_model=DetailedFeedbackResponse)
async def submit_detailed_feedback(feedback_data: DetailedFeedbackCreate):
    """
    Submit a detailed feedback form response.
    This endpoint is public (no auth required) so trainees can submit via the
    emailed link without needing to log in.
    """
    db = get_db()

    # Validate batch exists
    batch_result = db.table("batches").select("id, batch_name, end_date").eq("id", feedback_data.batchId).execute()
    if not batch_result.data:
        raise HTTPException(status_code=404, detail="Batch not found")

    batch = batch_result.data[0]
    end_date_str = batch.get("end_date", "")
    window_open, opens_on, closes_on = _get_feedback_window(end_date_str)
    if not window_open:
        now = datetime.utcnow()
        if opens_on and now < opens_on:
            raise HTTPException(
                status_code=400,
                detail=f"Feedback window has not opened yet. It opens on {opens_on.strftime('%d %b %Y')}."
            )
        else:
            raise HTTPException(status_code=400, detail="Feedback window has closed.")

    try:
        row = {
            "batch_id": feedback_data.batchId,
            "candidate_id": feedback_data.candidateId,
            "respondent_name": feedback_data.respondentName,
            "respondent_email": feedback_data.respondentEmail,
            "batch_no_and_trainer": feedback_data.batchNoAndTrainer,
            "takeaway1": feedback_data.takeaway1,
            "takeaway2": feedback_data.takeaway2,
            "takeaway3": feedback_data.takeaway3,
            "improvements": feedback_data.improvements,
            "course_impact": feedback_data.courseImpact,
            "trainer_rating": feedback_data.trainerRating,
            "assignments_helpful": feedback_data.assignmentsHelpful,
            "demonstrations_helpful": feedback_data.demonstrationsHelpful,
            "trainer_support_adequate": feedback_data.trainerSupportAdequate,
            "technical_discussions_helpful": feedback_data.technicalDiscussionsHelpful,
            "other_comments": feedback_data.otherComments,
            "submitted_at": datetime.utcnow().isoformat(),
        }
        result = db.table("detailed_feedbacks").insert(row).execute()
        if not result.data:
            raise HTTPException(status_code=500, detail="Failed to save feedback")

        saved = result.data[0]
        return DetailedFeedbackResponse(
            id=saved.get("id", ""),
            batchId=saved.get("batch_id", ""),
            candidateId=saved.get("candidate_id"),
            respondentName=saved.get("respondent_name"),
            respondentEmail=saved.get("respondent_email"),
            batchNoAndTrainer=saved.get("batch_no_and_trainer"),
            takeaway1=saved.get("takeaway1"),
            takeaway2=saved.get("takeaway2"),
            takeaway3=saved.get("takeaway3"),
            improvements=saved.get("improvements"),
            courseImpact=saved.get("course_impact"),
            trainerRating=saved.get("trainer_rating", 1),
            assignmentsHelpful=saved.get("assignments_helpful"),
            demonstrationsHelpful=saved.get("demonstrations_helpful"),
            trainerSupportAdequate=saved.get("trainer_support_adequate"),
            technicalDiscussionsHelpful=saved.get("technical_discussions_helpful"),
            otherComments=saved.get("other_comments"),
            submittedAt=datetime.fromisoformat(saved.get("submitted_at", datetime.utcnow().isoformat())),
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


# ─── Get Detailed Feedback for a Batch ──────────────────────────────────────

@router.get("/feedback/detailed/{batch_id}")
async def get_detailed_feedback(
    batch_id: str,
    current_user: dict = Depends(has_role("ADMIN", "COORDINATOR", "TRAINER"))
):
    """Get all detailed feedback responses for a batch."""
    db = get_db()
    check_batch_access(db, current_user, batch_id)

    try:
        result = db.table("detailed_feedbacks").select("*").eq("batch_id", batch_id).order("submitted_at").execute()
        return result.data or []
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


# ─── Export Detailed Feedback as Excel ──────────────────────────────────────

@router.get("/feedback/detailed/{batch_id}/export")
async def export_detailed_feedback_excel(
    batch_id: str,
    current_user: dict = Depends(has_role("ADMIN", "COORDINATOR", "TRAINER"))
):
    """
    Download all detailed feedback responses for a batch as an Excel file.
    The sheet format matches the Microsoft Forms export layout.
    """
    import io
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    db = get_db()
    check_batch_access(db, current_user, batch_id)

    batch_result = db.table("batches").select("batch_name").eq("id", batch_id).execute()
    batch_name = batch_result.data[0].get("batch_name", batch_id) if batch_result.data else batch_id

    result = db.table("detailed_feedbacks").select("*").eq("batch_id", batch_id).order("submitted_at").execute()
    rows = result.data or []

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feedback Responses"

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(border_style="thin", color="D3D3D3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    columns = [
        ("ID", 6),
        ("Start time", 18),
        ("Completion time", 18),
        ("Email", 28),
        ("Name", 22),
        ("Email Id", 28),
        ("Batch No and Trainer Name", 28),
        ("Top Takeaway 1", 40),
        ("Top Takeaway 2", 40),
        ("Top Takeaway 3", 40),
        ("What could have been done better?", 40),
        ("Impact of this course on you?", 40),
        ("Trainer Rating (1–5)", 14),
        ("Assignments helpful?", 20),
        ("Demonstrations helpful?", 20),
        ("Trainer support adequate?", 20),
        ("Technical discussions helpful?", 20),
        ("Other Comments", 40),
    ]

    for col_idx, (header_text, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 36

    alt_fill = PatternFill(start_color="EBF5FB", end_color="EBF5FB", fill_type="solid")

    for r_idx, row in enumerate(rows, 2):
        fill = alt_fill if r_idx % 2 == 0 else None
        submitted_at = row.get("submitted_at", "")

        values = [
            r_idx - 1,
            submitted_at,
            submitted_at,
            "anonymous",
            row.get("respondent_name", ""),
            row.get("respondent_email", ""),
            row.get("batch_no_and_trainer", ""),
            row.get("takeaway1", ""),
            row.get("takeaway2", ""),
            row.get("takeaway3", ""),
            row.get("improvements", ""),
            row.get("course_impact", ""),
            row.get("trainer_rating", ""),
            row.get("assignments_helpful", ""),
            row.get("demonstrations_helpful", ""),
            row.get("trainer_support_adequate", ""),
            row.get("technical_discussions_helpful", ""),
            row.get("other_comments", ""),
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = left if col_idx > 4 else center
            cell.border = border
            if fill:
                cell.fill = fill

        ws.row_dimensions[r_idx].height = 20

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    safe_name = batch_name.replace(" ", "_")
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="Feedback_{safe_name}.xlsx"'}
    )


# ─── Consolidated Reports and Downloads ──────────────────────────────────────

@router.get("/attendance-consolidated")
async def export_attendance_consolidated(
    batch_ids: str,
    pool_date: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    db = get_db()
    role = current_user.get("role")
    if role not in ["TRAINER", "COORDINATOR", "ADMIN"]:
        raise HTTPException(status_code=403, detail="Not authorized to access reports")
        
    ids_list = [bid.strip() for bid in batch_ids.split(",") if bid.strip()]
    for bid in ids_list:
        check_batch_access(db, current_user, bid)
        
    try:
        import io
        import openpyxl
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
        
        # 1. Fetch batches
        batches_res = db.table("batches").select("*").in_("id", ids_list).execute()
        batches_map = {b["id"]: b for b in (batches_res.data or [])}
        
        # 2. Fetch candidates for all selected batches
        cand_res = db.table("candidates").select("*").in_("batch_id", ids_list).execute()
        candidates = cand_res.data or []
        
        # Filter by pool date if provided
        emails_list = [c["email"].strip().lower() for c in candidates if c.get("email")]
        pool_map = {}
        if emails_list:
            pool_res = db.table("trainee_pool").select("email, onboarding_date").in_("email", list(set(emails_list))).execute()
            pool_map = {row["email"].strip().lower(): row.get("onboarding_date") for row in (pool_res.data or [])}
            
            if pool_date:
                candidates = [c for c in candidates if pool_map.get(c.get("email", "").strip().lower()) == pool_date]
        
        candidates.sort(key=lambda x: (x.get("batch_id", ""), x.get("full_name", "").lower()))
        
        # Fetch attendances for selected batches
        att_res = db.table("attendances").select("*").in_("batch_id", ids_list).execute()
        attendances = att_res.data or []
        
        # Group attendances by (candidate_id, batch_id)
        from collections import defaultdict
        att_groups = defaultdict(list)
        for att in attendances:
            att_groups[(att["candidate_id"], att["batch_id"])].append(att)
            
        wb = Workbook()
        ws = wb.active
        ws.title = "Consolidated Attendance"
        ws.views.sheetView[0].showGridLines = True
        
        # Styling
        title_font = Font(name="Calibri", size=16, bold=True, color="1F497D")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=11)
        bold_font = Font(name="Calibri", size=11, bold=True)
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center")
        thin_side = Side(border_style="thin", color="D3D3D3")
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        alt_fill = PatternFill(start_color="F2F4F4", end_color="F2F4F4", fill_type="solid")
        
        # Title
        ws.merge_cells("A1:J2")
        title_cell = ws["A1"]
        title_cell.value = "Consolidated Attendance Report"
        title_cell.font = title_font
        title_cell.alignment = center_align
        
        headers = [
            ("S.No", 6),
            ("Registration Number", 18),
            ("Name", 24),
            ("Email", 26),
            ("Onboarding Pool (Date)", 18),
            ("Batch Name", 24),
            ("Total Sessions", 14),
            ("Days Present", 12),
            ("Days Absent", 12),
            ("Attendance %", 14)
        ]
        
        for col_idx, (text, width) in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = border
            ws.column_dimensions[get_column_letter(col_idx)].width = width
            
        ws.row_dimensions[4].height = 28
        
        for r_idx, cand in enumerate(candidates, 5):
            cand_id = cand["id"]
            batch_id = cand["batch_id"]
            batch_name = batches_map.get(batch_id, {}).get("batch_name", "Unknown")
            email = cand.get("email", "").strip().lower()
            pool_date_val = pool_map.get(email, "N/A")
            
            cand_att = att_groups[(cand_id, batch_id)]
            total = len(cand_att)
            present = sum(1 for a in cand_att if a["status"] == "PRESENT")
            absent = sum(1 for a in cand_att if a["status"] == "ABSENT")
            
            pct = (present / total * 100) if total > 0 else 100.0
            
            fill = alt_fill if r_idx % 2 == 0 else None
            
            row_data = [
                r_idx - 4,
                cand.get("registration_number", "N/A"),
                cand.get("full_name", ""),
                cand.get("email", ""),
                pool_date_val,
                batch_name,
                total,
                present,
                absent,
                pct / 100.0
            ]
            
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=r_idx, column=col_idx, value=val)
                cell.font = regular_font if col_idx != 10 else bold_font
                cell.border = border
                if fill:
                    cell.fill = fill
                    
                if col_idx in [1, 5, 7, 8, 9]:
                    cell.alignment = center_align
                elif col_idx == 10:
                    cell.alignment = center_align
                    cell.number_format = '0.0%'
                else:
                    cell.alignment = left_align
            
            ws.row_dimensions[r_idx].height = 20
            
        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        
        return StreamingResponse(
            file_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="Consolidated_Attendance.xlsx"'}
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=400, detail=str(e))

@router.get("/assessment-scores")
async def export_assessment_scores(
    batch_id: str,
    current_user: dict = Depends(get_current_user)
):
    db = get_db()
    check_batch_access(db, current_user, batch_id)
    
    try:
        import io
        import openpyxl
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
        
        # Retrieve batch details
        batch_res = db.table("batches").select("batch_name").eq("id", batch_id).execute()
        batch_name = batch_res.data[0]["batch_name"] if batch_res.data else "Batch"
        
        # Query spark 1 and spark 2 report cards
        s1_res = db.table("spark_1_report_cards").select("*").eq("batch_id", batch_id).execute()
        s2_res = db.table("spark_2_report_cards").select("*").eq("batch_id", batch_id).execute()
        
        wb = Workbook()
        # Remove default sheet, we will add sheets based on what is available
        default_sheet = wb.active
        
        title_font = Font(name="Calibri", size=16, bold=True, color="1F497D")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=11)
        bold_font = Font(name="Calibri", size=11, bold=True)
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center")
        thin_side = Side(border_style="thin", color="D3D3D3")
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        alt_fill = PatternFill(start_color="F2F4F4", end_color="F2F4F4", fill_type="solid")
        
        def write_spark_sheet(ws, records, title):
            ws.views.sheetView[0].showGridLines = True
            ws.merge_cells("A1:Q2")
            title_cell = ws["A1"]
            title_cell.value = f"{title} - {batch_name}"
            title_cell.font = title_font
            title_cell.alignment = center_align
            
            headers = [
                ("Sr no", 6),
                ("Superset ID", 14),
                ("Name", 24),
                ("Email ID", 26),
                ("College", 24),
                ("Trainer Name", 20),
                ("Batch No", 12),
                ("A1 Score", 10),
                ("A2 Score", 10),
                ("Communication Skills", 12),
                ("Interpersonal Skills", 12),
                ("Business Etiquette", 12),
                ("Service Orientation", 12),
                ("Emotional Intelligence & Empathy", 14),
                ("Accountability & Ownership", 14),
                ("Presentation Skills", 12),
                ("Final Status", 14)
            ]
            
            for col_idx, (text, width) in enumerate(headers, 1):
                cell = ws.cell(row=4, column=col_idx, value=text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = border
                ws.column_dimensions[get_column_letter(col_idx)].width = width
                
            ws.row_dimensions[4].height = 28
            
            for r_idx, r in enumerate(records, 5):
                fill = alt_fill if r_idx % 2 == 0 else None
                row_data = [
                    r_idx - 4,
                    r.get("superset_id", "N/A"),
                    r.get("name", ""),
                    r.get("email", ""),
                    r.get("college", ""),
                    r.get("trainer_name", ""),
                    r.get("batch_no", ""),
                    r.get("a1_score"),
                    r.get("a2_score"),
                    r.get("communication_skills"),
                    r.get("interpersonal_skills"),
                    r.get("business_etiquette"),
                    r.get("service_orientation"),
                    r.get("emotional_intelligence_empathy"),
                    r.get("accountability_ownership"),
                    r.get("presentation_skills"),
                    r.get("final_status", "Not Cleared")
                ]
                
                for col_idx, val in enumerate(row_data, 1):
                    cell = ws.cell(row=r_idx, column=col_idx, value=val)
                    cell.font = regular_font if col_idx != 17 else bold_font
                    cell.border = border
                    if fill:
                        cell.fill = fill
                        
                    if col_idx in [1, 2, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17]:
                        cell.alignment = center_align
                    else:
                        cell.alignment = left_align
                ws.row_dimensions[r_idx].height = 20
                
        has_sheets = False
        if s1_res.data:
            ws1 = wb.create_sheet(title="Spark Phase 1")
            write_spark_sheet(ws1, s1_res.data, "Spark Phase 1 Report")
            has_sheets = True
            
        if s2_res.data:
            ws2 = wb.create_sheet(title="Spark Phase 2")
            write_spark_sheet(ws2, s2_res.data, "Spark Phase 2 Report")
            has_sheets = True
            
        if not has_sheets:
            # Fallback to candidate list and template sheet
            ws = wb.create_sheet(title="Spark Assessment Scores")
            cand_res = db.table("candidates").select("*").eq("batch_id", batch_id).execute()
            records = []
            for c in (cand_res.data or []):
                records.append({
                    "name": c.get("full_name"),
                    "email": c.get("email"),
                    "superset_id": c.get("registration_number"),
                    "college": c.get("college", "N/A"),
                    "trainer_name": "N/A",
                    "batch_no": batch_name,
                    "final_status": "Not Cleared"
                })
            write_spark_sheet(ws, records, "Spark Assessment Report Template")
            
        if default_sheet in wb.worksheets and len(wb.worksheets) > 1:
            wb.remove(default_sheet)
            
        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        
        safe_name = batch_name.replace(" ", "_")
        return StreamingResponse(
            file_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="Spark_Assessment_Scores_{safe_name}.xlsx"'}
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=400, detail=str(e))

@router.get("/consolidated")
async def export_consolidated_report(
    batch_ids: str,
    current_user: dict = Depends(get_current_user)
):
    db = get_db()
    role = current_user.get("role")
    if role not in ["TRAINER", "COORDINATOR", "ADMIN"]:
        raise HTTPException(status_code=403, detail="Not authorized to access reports")
        
    ids_list = [bid.strip() for bid in batch_ids.split(",") if bid.strip()]
    for bid in ids_list:
        check_batch_access(db, current_user, bid)
        
    try:
        import io
        import openpyxl
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
        
        # 1. Fetch batches
        batches_res = db.table("batches").select("*").in_("id", ids_list).execute()
        batches_map = {b["id"]: b for b in (batches_res.data or [])}
        
        # 2. Fetch candidates for all selected batches
        cand_res = db.table("candidates").select("*").in_("batch_id", ids_list).execute()
        candidates = cand_res.data or []
        
        emails_list = [c["email"].strip().lower() for c in candidates if c.get("email")]
        
        # Index pool info
        pool_map = {}
        if emails_list:
            pool_res = db.table("trainee_pool").select("*").in_("email", list(set(emails_list))).execute()
            pool_map = {row["email"].strip().lower(): row for row in (pool_res.data or [])}
            
        candidates.sort(key=lambda x: (x.get("batch_id", ""), x.get("full_name", "").lower()))
        
        # Fetch report cards for all candidates by email
        spark1_res = db.table("spark_1_report_cards").select("*").in_("email", emails_list).execute() if emails_list else None
        spark2_res = db.table("spark_2_report_cards").select("*").in_("email", emails_list).execute() if emails_list else None
        found_res = db.table("foundation_report_cards").select("*").in_("email", emails_list).execute() if emails_list else None
        stream_res = db.table("stream_report_cards").select("*").in_("email", emails_list).execute() if emails_list else None
        
        spark1_map = {r["email"].strip().lower(): r for r in (spark1_res.data or [])} if spark1_res else {}
        spark2_map = {r["email"].strip().lower(): r for r in (spark2_res.data or [])} if spark2_res else {}
        found_map = {r["email"].strip().lower(): r for r in (found_res.data or [])} if found_res else {}
        stream_map = {r["email"].strip().lower(): r for r in (stream_res.data or [])} if stream_res else {}
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Consolidated Report"
        ws.views.sheetView[0].showGridLines = True
        
        # Styling
        title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center")
        thin_side = Side(border_style="thin", color="D3D3D3")
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        title_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        
        # Headers definitions (Section headers on Row 3, detail headers on Row 4)
        ws.merge_cells("A1:AE2")
        title_cell = ws["A1"]
        title_cell.value = "Trainee Journey Consolidated Batch Report"
        title_cell.font = title_font
        title_cell.alignment = center_align
        title_cell.fill = title_fill
        
        sections = [
            ("Trainee Profile", 1, 7, "2E4053"),
            ("Spark Phase 1", 8, 13, "2E86C1"),
            ("Spark Phase 2", 14, 19, "28B463"),
            ("Foundational Training", 20, 24, "D35400"),
            ("Stream Based Training", 25, 31, "8E44AD")
        ]
        
        for name, start, end, color in sections:
            start_col = get_column_letter(start)
            end_col = get_column_letter(end)
            ws.merge_cells(f"{start_col}3:{end_col}3")
            cell = ws[f"{start_col}3"]
            cell.value = name
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.alignment = center_align
            fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            
            for c in range(start, end + 1):
                ws.cell(row=3, column=c).fill = fill
                ws.cell(row=3, column=c).border = border
                
        ws.row_dimensions[3].height = 24
        
        columns = [
            # Profile (1-7)
            ("S.No", 6),
            ("Reg Number", 14),
            ("Name", 22),
            ("Email", 24),
            ("College", 22),
            ("Onboarding Date", 16),
            ("Current Batch", 22),
            # Spark 1 (8-13)
            ("A1 Score", 10),
            ("A2 Score", 10),
            ("Soft Skills Avg", 12),
            ("Attendance %", 12),
            ("Status", 12),
            ("Trainer", 16),
            # Spark 2 (14-19)
            ("A1 Score", 10),
            ("A2 Score", 10),
            ("Soft Skills Avg", 12),
            ("Attendance %", 12),
            ("Status", 12),
            ("Trainer", 16),
            # Foundational (20-24)
            ("Language", 12),
            ("GA Avg", 10),
            ("Project Eval", 12),
            ("Final Grade", 12),
            ("Status", 12),
            # Stream (25-31)
            ("Stream Training", 22),
            ("MCQ Avg", 10),
            ("Coding Avg", 12),
            ("Project Avg", 12),
            ("Online Coding", 12),
            ("Attendance %", 12),
            ("Status", 12)
        ]
        
        for col_idx, (text, width) in enumerate(columns, 1):
            cell = ws.cell(row=4, column=col_idx, value=text)
            cell.font = Font(name="Calibri", size=9, bold=True, color="000000")
            cell.alignment = center_align
            cell.border = border
            ws.column_dimensions[get_column_letter(col_idx)].width = width
            cell.fill = PatternFill(start_color="EAEDED", end_color="EAEDED", fill_type="solid")
            
        ws.row_dimensions[4].height = 20
        
        alt_fill = PatternFill(start_color="F2F4F4", end_color="F2F4F4", fill_type="solid")
        
        for r_idx, cand in enumerate(candidates, 5):
            email = cand.get("email", "").strip().lower()
            batch_id = cand["batch_id"]
            batch_name = batches_map.get(batch_id, {}).get("batch_name", "Unknown")
            pool_info = pool_map.get(email, {})
            
            s1 = spark1_map.get(email, {})
            s2 = spark2_map.get(email, {})
            f = found_map.get(email, {})
            st = stream_map.get(email, {})
            
            def soft_avg(c):
                skills = [c.get("communication_skills"), c.get("interpersonal_skills"), c.get("business_etiquette"), c.get("service_orientation"), c.get("emotional_intelligence_empathy"), c.get("accountability_ownership"), c.get("presentation_skills")]
                valid = [s for s in skills if s is not None]
                return round(sum(valid) / len(valid) * 20.0, 1) if valid else "N/A"
                
            def found_ga_avg(c):
                gas_resolved = [TopperService.resolve_attempt(c.get(f"ga{i}_a1"), c.get(f"ga{i}_a2"), 60.0) for i in range(1, 6)]
                valid = [g for g in gas_resolved if g is not None]
                return round(sum(valid) / len(valid), 1) if valid else "N/A"
                
            def st_mcq_avg(c):
                mcqs = [TopperService.resolve_attempt(c.get(f"mcq{i}_a1"), c.get(f"mcq{i}_a2"), 60.0) for i in range(1, 8)]
                valid = [m for m in mcqs if m is not None]
                return round(sum(valid) / len(valid), 1) if valid else "N/A"
                
            def st_coding_avg(c):
                codings = [TopperService.resolve_attempt(c.get(f"coding{i}_a1"), c.get(f"coding{i}_a2"), 60.0) for i in range(1, 8)]
                valid = [co for co in codings if co is not None]
                return round(sum(valid) / len(valid), 1) if valid else "N/A"
                
            def st_proj_avg(c):
                projs = [TopperService.resolve_attempt(c.get(f"project_score{i}_a1"), c.get(f"project_score{i}_a2"), 65.0) for i in [1, 2]]
                valid = [p for p in projs if p is not None]
                return round(sum(valid) / len(valid), 1) if valid else "N/A"
            
            fill = alt_fill if r_idx % 2 == 0 else None
            
            row_data = [
                # Profile (1-7)
                r_idx - 4,
                cand.get("registration_number", "N/A"),
                cand.get("full_name", ""),
                cand.get("email", ""),
                cand.get("college", "N/A"),
                pool_info.get("onboarding_date", "N/A"),
                batch_name,
                # Spark 1 (8-13)
                s1.get("a1_score", "N/A"),
                s1.get("a2_score", "N/A"),
                soft_avg(s1),
                f"{s1.get('attendance_percentage', 0):.1f}%" if s1.get("attendance_percentage") is not None else "N/A",
                s1.get("final_status", "N/A"),
                s1.get("trainer_name", "N/A"),
                # Spark 2 (14-19)
                s2.get("a1_score", "N/A"),
                s2.get("a2_score", "N/A"),
                soft_avg(s2),
                f"{s2.get('attendance_percentage', 0):.1f}%" if s2.get("attendance_percentage") is not None else "N/A",
                s2.get("final_status", "N/A"),
                s2.get("trainer_name", "N/A"),
                # Foundational (20-24)
                f.get("foundation_language", "N/A"),
                found_ga_avg(f),
                TopperService.resolve_attempt(f.get("project_eval_a1"), f.get("project_eval_a2"), 65.0) if f.get("project_eval_a1") is not None else "N/A",
                TopperService.resolve_attempt(f.get("final_grade_a1"), f.get("final_grade_a2"), 60.0) if f.get("final_grade_a1") is not None else "N/A",
                f.get("training_status", "N/A"),
                # Stream (25-31)
                st.get("stream_training", "N/A"),
                st_mcq_avg(st),
                st_coding_avg(st),
                st_proj_avg(st),
                TopperService.resolve_attempt(st.get("online_coding_a1"), st.get("online_coding_a2"), 60.0) if st.get("online_coding_a1") is not None else "N/A",
                f"{st.get('attendance_percentage', 0):.1f}%" if st.get("attendance_percentage") is not None else "N/A",
                st.get("final_status", "N/A")
            ]
            
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=r_idx, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = border
                if fill:
                    cell.fill = fill
                
                if col_idx in [1, 2, 6, 8, 9, 10, 11, 12, 14, 15, 16, 17, 18, 20, 21, 22, 23, 24, 26, 27, 28, 29, 30, 31]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align
            ws.row_dimensions[r_idx].height = 20
            
        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        
        return StreamingResponse(
            file_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="Consolidated_Batch_Report.xlsx"'}
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=400, detail=str(e))



